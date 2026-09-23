# ruff: noqa: B008, E501, E701, E702, B904
import json
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Annotated, Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.models import (
    AuditLog,
    CertificateCase,
    CertificateDeliveryEvent,
    CourseCatalog,
    CourseEnrollment,
    CourseOfferingVersion,
    CourseSubjectVersion,
    ExamAttempt,
    ExamBatch,
    ExamBatchSubject,
    ExamDomainAttachment,
    ExamFeeAssessment,
    ExamRegistration,
    ExamRegistrationSubject,
    ExamResultImportBatch,
    ExamResultImportLine,
    ExamResultRevision,
    ExamSchemeVersion,
    FeePolicyVersion,
    FileObject,
    FileReplica,
    Student,
    User,
    utc_now,
)
from app.core.pii import decrypt, encrypt
from app.modules.attachments.router import clean_filename, detect, extension
from app.modules.attachments.storage import StorageError, get_storage
from app.modules.iam.router import Db, require_permission

router = APIRouter(prefix="/api", tags=["exam"])
BatchRead = Annotated[User, Depends(require_permission("exam.batch.read"))]
BatchManage = Annotated[User, Depends(require_permission("exam.batch.manage"))]
RegistrationManage = Annotated[User, Depends(require_permission("exam.registration.manage"))]
RegistrationRead = Annotated[User, Depends(require_permission("exam.registration.read"))]
ResultRead = Annotated[User, Depends(require_permission("exam.result.read"))]
ResultEdit = Annotated[User, Depends(require_permission("exam.result.edit"))]
ResultSubmit = Annotated[User, Depends(require_permission("exam.result.submit"))]
ResultImport = Annotated[User, Depends(require_permission("exam.result.import"))]
ResultConfirm = Annotated[User, Depends(require_permission("exam.result.confirm"))]
ResultRevise = Annotated[User, Depends(require_permission("exam.result.revise"))]
FeeRead = Annotated[User, Depends(require_permission("exam_fee_assessment.read"))]
FeeConfirm = Annotated[User, Depends(require_permission("exam_fee_assessment.confirm"))]
FeeWaive = Annotated[User, Depends(require_permission("exam_fee_assessment.waive"))]
CertificateRead = Annotated[User, Depends(require_permission("certificate.read"))]
CertificateManage = Annotated[User, Depends(require_permission("certificate.manage"))]
CertificateDeliver = Annotated[User, Depends(require_permission("certificate.deliver"))]
CertificateSensitive = Annotated[User, Depends(require_permission("certificate.sensitive.reveal"))]
CertificateAttachmentRead = Annotated[
    User, Depends(require_permission("certificate.attachment.read"))
]


def error(status: int, code: str, message: str) -> HTTPException:
    return HTTPException(status_code=status, detail={"code": code, "message": message})


def new_code(prefix: str) -> str:
    return f"{prefix}{uuid.uuid4().hex[:12].upper()}"


def audit(
    db: Session,
    user: User,
    request: Request,
    action: str,
    kind: str,
    item_id: str,
    detail: dict[str, Any] | None = None,
) -> None:
    db.add(
        AuditLog(
            organization_id=user.organization_id,
            actor_user_id=user.id,
            action=action,
            subject_type=kind,
            subject_id=item_id,
            correlation_id=request.state.correlation_id,
            detail_json=json.dumps(detail or {}, ensure_ascii=False),
        )
    )


def get_batch_for_user(db: Session, user: User, batch_id: str) -> ExamBatch:
    item = db.scalar(
        select(ExamBatch).where(
            ExamBatch.id == batch_id,
            ExamBatch.organization_id == user.organization_id,
        )
    )
    if not item:
        raise error(404, "EXAM_BATCH_NOT_FOUND", "未找到考试批次。")
    return item


def get_attempt_for_user(db: Session, user: User, attempt_id: str) -> ExamAttempt:
    item = db.scalar(
        select(ExamAttempt)
        .join(CourseEnrollment, ExamAttempt.course_enrollment_id == CourseEnrollment.id)
        .where(
            ExamAttempt.id == attempt_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not item:
        raise error(404, "EXAM_ATTEMPT_NOT_FOUND", "未找到考试尝试。")
    return item


def snapshot(item: ExamAttempt) -> dict[str, Any]:
    return {
        "score_value_scaled": item.score_value_scaled,
        "score_scale": item.score_scale,
        "result_status": item.result_status,
        "attendance_status": item.attendance_status,
        "result_confirm_status": item.result_confirm_status,
    }


def attempt_view(db: Session, item: ExamAttempt) -> dict[str, Any]:
    subject = db.get(CourseSubjectVersion, item.course_subject_version_id)
    attachments = db.execute(
        select(ExamDomainAttachment, FileObject)
        .join(FileObject, ExamDomainAttachment.file_object_id == FileObject.id)
        .where(
            ExamDomainAttachment.owner_type == "EXAM_ATTEMPT",
            ExamDomainAttachment.owner_id == item.id,
        )
    ).all()
    return {
        "id": item.id,
        "course_enrollment_id": item.course_enrollment_id,
        "subject": {"id": subject.id, "code": subject.subject_code, "name": subject.subject_name},
        "attempt_no": item.attempt_no,
        "attempt_type": item.attempt_type,
        "score_value_scaled": item.score_value_scaled,
        "score_scale": item.score_scale,
        "result_status": item.result_status,
        "result_confirm_status": item.result_confirm_status,
        "attendance_status": item.attendance_status,
        "current_revision_no": item.current_revision_no,
        "attachments": [
            {
                "id": link.id,
                "file_object_id": file.id,
                "attachment_type": link.attachment_type,
                "original_filename": file.original_filename,
            }
            for link, file in attachments
        ],
        "version": item.version,
    }


def recompute(db: Session, enrollment: CourseEnrollment) -> str:
    subjects = db.scalars(
        select(CourseSubjectVersion).where(
            CourseSubjectVersion.exam_scheme_version_id == enrollment.exam_scheme_version_id,
            CourseSubjectVersion.enabled.is_(True),
        )
    ).all()
    passed = set(
        db.scalars(
            select(ExamAttempt.course_subject_version_id).where(
                ExamAttempt.course_enrollment_id == enrollment.id,
                ExamAttempt.result_status == "PASSED",
                ExamAttempt.result_confirm_status == "OFFICIALLY_CONFIRMED",
            )
        ).all()
    )
    pending = (
        db.scalar(
            select(func.count())
            .select_from(ExamAttempt)
            .where(
                ExamAttempt.course_enrollment_id == enrollment.id,
                ExamAttempt.result_confirm_status != "OFFICIALLY_CONFIRMED",
            )
        )
        or 0
    )
    attempts = (
        db.scalar(
            select(func.count())
            .select_from(ExamAttempt)
            .where(ExamAttempt.course_enrollment_id == enrollment.id)
        )
        or 0
    )
    if subjects and all(s.id in passed for s in subjects):
        status = "PASSED"
    elif passed:
        status = "PARTIALLY_PASSED"
    elif pending:
        status = "IN_PROGRESS"
    elif attempts:
        status = "FAILED"
    else:
        status = "NOT_REGISTERED"
    if enrollment.exam_status != status:
        enrollment.exam_status = status
        enrollment.version += 1
    certificate = db.scalar(
        select(CertificateCase).where(CertificateCase.course_enrollment_id == enrollment.id)
    )
    if not certificate:
        enrollment.certificate_status = "ELIGIBLE" if status == "PASSED" else "NOT_ELIGIBLE"
    return status


class BatchInput(BaseModel):
    batch_code: str = Field(min_length=1, max_length=40)
    batch_name: str = Field(min_length=1, max_length=160)
    exam_scheme_version_id: str
    exam_start_date: date
    exam_end_date: date
    registration_deadline_at: datetime | None = None
    organizing_institution: str | None = None


class SubjectInput(BaseModel):
    course_subject_version_id: str
    exam_start_at: datetime
    exam_end_at: datetime | None = None
    venue_name: str | None = None
    capacity: int | None = Field(default=None, gt=0)
    initial_exam_fee_amount_cent: int | None = Field(default=None, ge=0)
    resit_fee_amount_cent: int | None = Field(default=None, ge=0)


class VersionInput(BaseModel):
    version: int


class ReasonInput(VersionInput):
    reason: str = Field(min_length=1, max_length=500)


class RegistrationInput(BaseModel):
    course_enrollment_id: str
    exam_batch_id: str
    exam_batch_subject_ids: list[str] = Field(min_length=1)
    idempotency_key: str = Field(min_length=8, max_length=100)


class ResitRegistrationInput(BaseModel):
    exam_batch_id: str
    exam_batch_subject_ids: list[str] = Field(min_length=1)
    idempotency_key: str = Field(min_length=8, max_length=100)


class AdmissionTicketInput(VersionInput):
    admission_ticket_no: str = Field(min_length=4, max_length=100)


class ResultInput(BaseModel):
    score_value_scaled: int | None = Field(default=None, ge=0)
    score_scale: int = Field(default=100, gt=0)
    result_status: str
    attendance_status: str = "PRESENT"
    remarks: str | None = None
    version: int


class RevisionInput(ResultInput):
    reason: str = Field(min_length=1, max_length=500)


class FeeAction(VersionInput):
    reason: str | None = Field(default=None, max_length=500)


class CertificateAction(VersionInput):
    remarks: str | None = Field(default=None, max_length=500)
    certificate_no: str | None = None
    issuing_authority: str | None = None
    delivery_method: str | None = None
    recipient_name: str | None = None
    courier_company: str | None = None
    tracking_no: str | None = None


class ImportConfirmInput(BaseModel):
    idempotency_key: str = Field(min_length=8, max_length=100)


class RevealInput(BaseModel):
    reason: str = Field(min_length=1, max_length=500)


@router.get("/exam-batches")
def batches(db: Db, _user: BatchRead) -> dict[str, Any]:
    return {
        "items": [
            {
                "id": x.id,
                "batch_code": x.batch_code,
                "batch_name": x.batch_name,
                "batch_status": x.batch_status,
                "exam_start_date": x.exam_start_date,
                "exam_end_date": x.exam_end_date,
                "version": x.version,
            }
            for x in db.scalars(
                select(ExamBatch)
                .where(ExamBatch.organization_id == _user.organization_id)
                .order_by(ExamBatch.exam_start_date.desc())
            ).all()
        ]
    }


@router.get("/exam-schemes")
def exam_schemes(db: Db, _user: BatchRead) -> dict[str, Any]:
    return {
        "items": [
            {
                "id": scheme.id,
                "version_no": scheme.version_no,
                "name": scheme.name,
                "subjects": [
                    {
                        "id": subject.id,
                        "subject_code": subject.subject_code,
                        "subject_name": subject.subject_name,
                    }
                    for subject in db.scalars(
                        select(CourseSubjectVersion)
                        .where(CourseSubjectVersion.exam_scheme_version_id == scheme.id)
                        .order_by(CourseSubjectVersion.display_order)
                    ).all()
                ],
            }
            for scheme in db.scalars(
                select(ExamSchemeVersion)
                .join(
                    CourseOfferingVersion,
                    CourseOfferingVersion.exam_scheme_version_id == ExamSchemeVersion.id,
                )
                .join(
                    CourseCatalog,
                    CourseOfferingVersion.course_catalog_id == CourseCatalog.id,
                )
                .where(CourseCatalog.organization_id == _user.organization_id)
                .distinct()
                .order_by(ExamSchemeVersion.created_at.desc())
            ).all()
        ]
    }


@router.post("/exam-batches", status_code=201)
def create_batch(data: BatchInput, request: Request, db: Db, user: BatchManage) -> dict[str, Any]:
    if data.exam_end_date < data.exam_start_date:
        raise error(422, "EXAM_DATE_INVALID", "结束日期不能早于开始日期。")
    if not db.get(ExamSchemeVersion, data.exam_scheme_version_id):
        raise error(422, "EXAM_SCHEME_NOT_FOUND", "考试方案不存在。")
    item = ExamBatch(
        **data.model_dump(),
        organization_id=user.organization_id,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(item)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise error(409, "EXAM_BATCH_CODE_EXISTS", "考试批次编码已存在。")
    audit(db, user, request, "exam.batch.create", "exam_batch", item.id)
    db.commit()
    return {"id": item.id, "batch_status": item.batch_status, "version": item.version}


@router.get("/exam-batches/{batch_id}")
def batch_detail(batch_id: str, db: Db, _user: BatchRead) -> dict[str, Any]:
    item = get_batch_for_user(db, _user, batch_id)
    subjects = db.execute(
        select(ExamBatchSubject, CourseSubjectVersion)
        .join(CourseSubjectVersion)
        .where(ExamBatchSubject.exam_batch_id == item.id)
    ).all()
    return {
        "id": item.id,
        "batch_code": item.batch_code,
        "batch_name": item.batch_name,
        "exam_scheme_version_id": item.exam_scheme_version_id,
        "batch_status": item.batch_status,
        "version": item.version,
        "subjects": [
            {
                "id": x.id,
                "subject_code": s.subject_code,
                "subject_name": s.subject_name,
                "exam_start_at": x.exam_start_at,
                "initial_exam_fee_amount_cent": x.initial_exam_fee_amount_cent,
                "resit_fee_amount_cent": x.resit_fee_amount_cent,
            }
            for x, s in subjects
        ],
    }


@router.post("/exam-batches/{batch_id}/subjects", status_code=201)
def add_subject(
    batch_id: str, data: SubjectInput, request: Request, db: Db, user: BatchManage
) -> dict[str, Any]:
    batch = get_batch_for_user(db, user, batch_id)
    if batch.batch_status != "DRAFT":
        raise error(409, "EXAM_BATCH_LOCKED", "只有草稿批次可以增加科目。")
    subject = db.get(CourseSubjectVersion, data.course_subject_version_id)
    if not subject or subject.exam_scheme_version_id != batch.exam_scheme_version_id:
        raise error(422, "EXAM_SUBJECT_SCHEME_MISMATCH", "科目不属于批次考试方案。")
    item = ExamBatchSubject(exam_batch_id=batch.id, **data.model_dump())
    db.add(item)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise error(409, "EXAM_BATCH_SUBJECT_EXISTS", "批次科目已存在。")
    audit(db, user, request, "exam.batch.subject.create", "exam_batch_subject", item.id)
    db.commit()
    return {"id": item.id}


def batch_transition(
    batch_id: str,
    data: VersionInput | ReasonInput,
    request: Request,
    db: Session,
    user: User,
    target: str,
) -> dict[str, Any]:
    item = get_batch_for_user(db, user, batch_id)
    allowed = {
        "OPEN": {"DRAFT"},
        "REGISTRATION_CLOSED": {"OPEN"},
        "IN_PROGRESS": {"REGISTRATION_CLOSED"},
        "COMPLETED": {"IN_PROGRESS", "RESULT_PENDING"},
        "CANCELLED": {"DRAFT", "OPEN", "REGISTRATION_CLOSED"},
    }
    if item.version != data.version or item.batch_status not in allowed[target]:
        raise error(409, "EXAM_BATCH_STATUS_CONFLICT", "批次状态不能执行该操作。")
    if target == "CANCELLED":
        item.cancellation_reason = data.reason
    item.batch_status = target
    item.version += 1
    audit(db, user, request, f"exam.batch.{target.lower()}", "exam_batch", item.id)
    db.commit()
    return {"id": item.id, "batch_status": item.batch_status, "version": item.version}


@router.post("/exam-batches/{batch_id}/open")
def open_batch(batch_id: str, data: VersionInput, request: Request, db: Db, user: BatchManage):
    return batch_transition(batch_id, data, request, db, user, "OPEN")


@router.post("/exam-batches/{batch_id}/close-registration")
def close_batch(batch_id: str, data: VersionInput, request: Request, db: Db, user: BatchManage):
    return batch_transition(batch_id, data, request, db, user, "REGISTRATION_CLOSED")


@router.post("/exam-batches/{batch_id}/start")
def start_batch(batch_id: str, data: VersionInput, request: Request, db: Db, user: BatchManage):
    return batch_transition(batch_id, data, request, db, user, "IN_PROGRESS")


@router.post("/exam-batches/{batch_id}/complete")
def complete_batch(batch_id: str, data: VersionInput, request: Request, db: Db, user: BatchManage):
    return batch_transition(batch_id, data, request, db, user, "COMPLETED")


@router.post("/exam-batches/{batch_id}/cancel")
def cancel_batch(batch_id: str, data: ReasonInput, request: Request, db: Db, user: BatchManage):
    return batch_transition(batch_id, data, request, db, user, "CANCELLED")


@router.post("/exam-registrations", status_code=201)
def register(
    data: RegistrationInput, request: Request, db: Db, user: RegistrationManage
) -> dict[str, Any]:
    old = db.scalar(
        select(ExamRegistration)
        .join(CourseEnrollment)
        .where(
            ExamRegistration.idempotency_key == data.idempotency_key,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if old:
        return {"id": old.id, "registration_no": old.registration_no, "idempotent": True}
    enrollment = db.get(CourseEnrollment, data.course_enrollment_id)
    batch = get_batch_for_user(db, user, data.exam_batch_id)
    if (
        not enrollment
        or enrollment.organization_id != user.organization_id
        or enrollment.exam_scheme_version_id != batch.exam_scheme_version_id
        or batch.batch_status != "OPEN"
    ):
        raise error(409, "EXAM_REGISTRATION_INVALID", "报名或考试批次状态不符合要求。")
    if batch.registration_deadline_at and batch.registration_deadline_at < utc_now():
        raise error(409, "EXAM_REGISTRATION_CLOSED", "考试报名截止时间已过。")
    selected = db.scalars(
        select(ExamBatchSubject).where(
            ExamBatchSubject.id.in_(data.exam_batch_subject_ids),
            ExamBatchSubject.exam_batch_id == batch.id,
        )
    ).all()
    if len(selected) != len(set(data.exam_batch_subject_ids)):
        raise error(422, "EXAM_SUBJECT_INVALID", "选择的批次科目无效。")
    registration = ExamRegistration(
        registration_no=new_code("ER"),
        course_enrollment_id=enrollment.id,
        exam_batch_id=batch.id,
        registration_source="INITIAL",
        idempotency_key=data.idempotency_key,
        registration_status="REGISTERED",
        registered_at=utc_now(),
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(registration)
    db.flush()
    attempt_types: set[str] = set()
    for batch_subject in selected:
        prior = db.scalars(
            select(ExamAttempt)
            .where(
                ExamAttempt.course_enrollment_id == enrollment.id,
                ExamAttempt.course_subject_version_id == batch_subject.course_subject_version_id,
            )
            .order_by(ExamAttempt.attempt_no)
        ).all()
        if any(
            x.result_status == "PASSED" and x.result_confirm_status == "OFFICIALLY_CONFIRMED"
            for x in prior
        ):
            raise error(409, "SUBJECT_ALREADY_PASSED", "已通过科目不能重复报名。")
        if prior and prior[-1].result_confirm_status != "OFFICIALLY_CONFIRMED":
            raise error(409, "PREVIOUS_RESULT_PENDING", "上一考试结果尚未确认。")
        attempt_no = len(prior) + 1
        attempt_type = "INITIAL" if attempt_no == 1 else "RESIT"
        attempt_types.add(attempt_type)
        rs = ExamRegistrationSubject(
            exam_registration_id=registration.id,
            exam_batch_subject_id=batch_subject.id,
            course_subject_version_id=batch_subject.course_subject_version_id,
            attempt_type=attempt_type,
            planned_attempt_no=attempt_no,
            subject_registration_status="REGISTERED",
            created_by=user.id,
        )
        db.add(rs)
        db.flush()
        attempt = ExamAttempt(
            course_enrollment_id=enrollment.id,
            course_subject_version_id=batch_subject.course_subject_version_id,
            exam_registration_subject_id=rs.id,
            attempt_no=attempt_no,
            attempt_type=attempt_type,
            created_by=user.id,
            updated_by=user.id,
        )
        db.add(attempt)
        db.flush()
        amount = (
            batch_subject.initial_exam_fee_amount_cent
            if attempt_type == "INITIAL"
            else batch_subject.resit_fee_amount_cent
        )
        if amount is None:
            policy = db.get(FeePolicyVersion, enrollment.fee_policy_version_id)
            amount = policy.initial_exam_fee_cent if attempt_type == "INITIAL" else None
        if amount is None:
            raise error(409, "EXAM_FEE_POLICY_MISSING", "无法确定考试费用，报名已阻止。")
        db.add(
            ExamFeeAssessment(
                assessment_no=new_code("EFA"),
                course_enrollment_id=enrollment.id,
                exam_attempt_id=attempt.id,
                exam_registration_subject_id=rs.id,
                course_subject_version_id=batch_subject.course_subject_version_id,
                fee_policy_version_id=enrollment.fee_policy_version_id,
                fee_type="INITIAL_EXAM" if attempt_type == "INITIAL" else "RESIT",
                responsibility="SCHOOL" if attempt_type == "INITIAL" else "STUDENT",
                amount_cent=amount,
                policy_snapshot_description=f"{attempt_type} fee snapshot",
                created_by=user.id,
                updated_by=user.id,
            )
        )
    registration.registration_source = (
        "MIXED" if len(attempt_types) > 1 else next(iter(attempt_types))
    )
    enrollment.exam_status = "REGISTERED"
    audit(db, user, request, "exam.registration.create", "exam_registration", registration.id)
    db.commit()
    return {
        "id": registration.id,
        "registration_no": registration.registration_no,
        "idempotent": False,
    }


def registration_view(db: Session, item: ExamRegistration) -> dict[str, Any]:
    batch = db.get(ExamBatch, item.exam_batch_id)
    rows = db.execute(
        select(ExamRegistrationSubject, CourseSubjectVersion, ExamAttempt)
        .join(
            CourseSubjectVersion,
            CourseSubjectVersion.id == ExamRegistrationSubject.course_subject_version_id,
        )
        .join(
            ExamAttempt,
            ExamAttempt.exam_registration_subject_id == ExamRegistrationSubject.id,
        )
        .where(ExamRegistrationSubject.exam_registration_id == item.id)
    ).all()
    return {
        "id": item.id,
        "registration_no": item.registration_no,
        "course_enrollment_id": item.course_enrollment_id,
        "exam_batch": {
            "id": batch.id,
            "code": batch.batch_code,
            "name": batch.batch_name,
        },
        "registration_status": item.registration_status,
        "registration_source": item.registration_source,
        "version": item.version,
        "subjects": [
            {
                "id": row.id,
                "subject_code": subject.subject_code,
                "subject_name": subject.subject_name,
                "attempt_type": row.attempt_type,
                "planned_attempt_no": row.planned_attempt_no,
                "admission_ticket_no_masked": (
                    f"****{row.admission_ticket_no_last4}"
                    if row.admission_ticket_no_last4
                    else None
                ),
                "attempt": attempt_view(db, attempt),
            }
            for row, subject, attempt in rows
        ],
    }


@router.get("/exam-registrations")
def registrations(
    db: Db,
    _user: RegistrationRead,
    enrollment_id: str | None = None,
    batch_id: str | None = None,
) -> dict[str, Any]:
    statement = select(ExamRegistration)
    statement = statement.join(
        CourseEnrollment,
        ExamRegistration.course_enrollment_id == CourseEnrollment.id,
    ).where(CourseEnrollment.organization_id == _user.organization_id)
    if enrollment_id:
        statement = statement.where(ExamRegistration.course_enrollment_id == enrollment_id)
    if batch_id:
        statement = statement.where(ExamRegistration.exam_batch_id == batch_id)
    return {
        "items": [
            registration_view(db, item)
            for item in db.scalars(statement.order_by(ExamRegistration.created_at.desc())).all()
        ]
    }


@router.get("/exam-registrations/{registration_id}")
def registration_detail(registration_id: str, db: Db, _user: RegistrationRead) -> dict[str, Any]:
    item = db.scalar(
        select(ExamRegistration)
        .join(CourseEnrollment)
        .where(
            ExamRegistration.id == registration_id,
            CourseEnrollment.organization_id == _user.organization_id,
        )
    )
    if not item:
        raise error(404, "EXAM_REGISTRATION_NOT_FOUND", "未找到考试报名。")
    return registration_view(db, item)


@router.post("/enrollments/{enrollment_id}/resit-register", status_code=201)
def resit_register(
    enrollment_id: str,
    data: ResitRegistrationInput,
    request: Request,
    db: Db,
    user: RegistrationManage,
) -> dict[str, Any]:
    return register(
        RegistrationInput(course_enrollment_id=enrollment_id, **data.model_dump()),
        request,
        db,
        user,
    )


@router.post("/exam-registration-subjects/{subject_id}/admission-ticket")
def set_admission_ticket(
    subject_id: str,
    data: AdmissionTicketInput,
    request: Request,
    db: Db,
    user: RegistrationManage,
) -> dict[str, Any]:
    item = db.scalar(
        select(ExamRegistrationSubject)
        .join(ExamRegistration)
        .join(CourseEnrollment)
        .where(
            ExamRegistrationSubject.id == subject_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not item:
        raise error(404, "EXAM_REGISTRATION_SUBJECT_NOT_FOUND", "未找到科目报名。")
    if item.version != data.version:
        raise error(409, "VERSION_CONFLICT", "科目报名已被其他操作修改。")
    item.admission_ticket_no_ciphertext = encrypt(data.admission_ticket_no)
    item.admission_ticket_no_last4 = data.admission_ticket_no[-4:]
    item.subject_registration_status = "ADMITTED"
    item.version += 1
    audit(db, user, request, "exam.admission_ticket.set", "exam_registration_subject", item.id)
    db.commit()
    return {
        "id": item.id,
        "admission_ticket_no_masked": f"****{item.admission_ticket_no_last4}",
        "version": item.version,
    }


@router.post("/enrollments/{enrollment_id}/resit-preview")
def resit_preview(enrollment_id: str, db: Db, _user: RegistrationRead) -> dict[str, Any]:
    progress = exam_progress(enrollment_id, db, _user)
    return {
        "enrollment_id": enrollment_id,
        "eligible_subjects": [
            item
            for item in progress["subjects"]
            if not item["passed"]
            and (
                not item["attempts"]
                or item["attempts"][-1]["result_confirm_status"] == "OFFICIALLY_CONFIRMED"
            )
        ],
        "blocked_subjects": [
            item
            for item in progress["subjects"]
            if item["passed"]
            or (
                item["attempts"]
                and item["attempts"][-1]["result_confirm_status"] != "OFFICIALLY_CONFIRMED"
            )
        ],
    }


@router.get("/enrollments/{enrollment_id}/exam-progress")
def exam_progress(enrollment_id: str, db: Db, _user: ResultRead) -> dict[str, Any]:
    enrollment = db.get(CourseEnrollment, enrollment_id)
    if not enrollment or enrollment.organization_id != _user.organization_id:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名。")
    subjects = db.scalars(
        select(CourseSubjectVersion)
        .where(
            CourseSubjectVersion.exam_scheme_version_id == enrollment.exam_scheme_version_id,
            CourseSubjectVersion.enabled.is_(True),
        )
        .order_by(CourseSubjectVersion.display_order)
    ).all()
    items = []
    for subject in subjects:
        attempts = db.scalars(
            select(ExamAttempt)
            .where(
                ExamAttempt.course_enrollment_id == enrollment.id,
                ExamAttempt.course_subject_version_id == subject.id,
            )
            .order_by(ExamAttempt.attempt_no)
        ).all()
        passed = next(
            (
                x
                for x in attempts
                if x.result_status == "PASSED" and x.result_confirm_status == "OFFICIALLY_CONFIRMED"
            ),
            None,
        )
        items.append(
            {
                "subject_id": subject.id,
                "subject_code": subject.subject_code,
                "subject_name": subject.subject_name,
                "attempt_count": len(attempts),
                "passed": bool(passed),
                "passed_attempt_no": passed.attempt_no if passed else None,
                "attempts": [attempt_view(db, x) for x in attempts],
            }
        )
    return {
        "enrollment_id": enrollment.id,
        "exam_status": enrollment.exam_status,
        "subjects": items,
    }


@router.get("/exam-attempts/{attempt_id}")
def attempt_detail(attempt_id: str, db: Db, _user: ResultRead) -> dict[str, Any]:
    return attempt_view(db, get_attempt_for_user(db, _user, attempt_id))


@router.patch("/exam-attempts/{attempt_id}/draft-result")
def draft_result(
    attempt_id: str, data: ResultInput, request: Request, db: Db, user: ResultEdit
) -> dict[str, Any]:
    item = get_attempt_for_user(db, user, attempt_id)
    if item.version != data.version or item.result_confirm_status != "DRAFT":
        raise error(409, "EXAM_RESULT_LOCKED", "成绩当前不可编辑。")
    if data.result_status not in {"PENDING", "PASSED", "FAILED", "ABSENT", "INVALID", "CANCELLED"}:
        raise error(422, "EXAM_RESULT_INVALID", "成绩状态无效。")
    subject = db.get(CourseSubjectVersion, item.course_subject_version_id)
    if data.score_value_scaled is not None and data.score_value_scaled > 100 * data.score_scale:
        raise error(422, "EXAM_SCORE_INVALID", "分数超出范围。")
    suggested = (
        "PASSED"
        if data.score_value_scaled is not None
        and subject.passing_score is not None
        and data.score_value_scaled >= subject.passing_score * data.score_scale
        else "FAILED"
    )
    if (
        data.score_value_scaled is not None
        and data.result_status in {"PASSED", "FAILED"}
        and data.result_status != suggested
    ):
        raise error(422, "EXAM_RESULT_SCORE_MISMATCH", "分数与结果不一致。")
    item.score_value_scaled = data.score_value_scaled
    item.score_scale = data.score_scale
    item.result_status = data.result_status
    item.attendance_status = data.attendance_status
    item.remarks = data.remarks
    item.version += 1
    item.updated_by = user.id
    audit(db, user, request, "exam.result.draft", "exam_attempt", item.id)
    db.commit()
    return attempt_view(db, item)


@router.post("/exam-attempts/{attempt_id}/submit-result")
def submit_result(
    attempt_id: str, data: VersionInput, request: Request, db: Db, user: ResultSubmit
) -> dict[str, Any]:
    item = get_attempt_for_user(db, user, attempt_id)
    if (
        item.version != data.version
        or item.result_confirm_status != "DRAFT"
        or item.result_status == "PENDING"
    ):
        raise error(409, "EXAM_RESULT_STATUS_CONFLICT", "成绩不能提交。")
    item.result_confirm_status = "SUBMITTED"
    item.version += 1
    audit(db, user, request, "exam.result.submit", "exam_attempt", item.id)
    db.commit()
    return attempt_view(db, item)


@router.post("/exam-attempts/{attempt_id}/confirm-result")
def confirm_result(
    attempt_id: str, data: VersionInput, request: Request, db: Db, user: ResultConfirm
) -> dict[str, Any]:
    item = get_attempt_for_user(db, user, attempt_id)
    if item.version != data.version or item.result_confirm_status != "SUBMITTED":
        raise error(409, "EXAM_RESULT_STATUS_CONFLICT", "成绩不能确认。")
    item.result_confirm_status = "OFFICIALLY_CONFIRMED"
    item.confirmed_at = utc_now()
    item.confirmed_by = user.id
    item.version += 1
    enrollment = db.get(CourseEnrollment, item.course_enrollment_id)
    status = recompute(db, enrollment)
    audit(
        db,
        user,
        request,
        "exam.result.confirm",
        "exam_attempt",
        item.id,
        {"enrollment_exam_status": status},
    )
    db.commit()
    return attempt_view(db, item)


@router.post("/exam-attempts/{attempt_id}/return-result")
def return_result(
    attempt_id: str,
    data: ReasonInput,
    request: Request,
    db: Db,
    user: ResultConfirm,
) -> dict[str, Any]:
    item = get_attempt_for_user(db, user, attempt_id)
    if item.version != data.version or item.result_confirm_status != "SUBMITTED":
        raise error(409, "EXAM_RESULT_STATUS_CONFLICT", "成绩不能退回。")
    item.result_confirm_status = "DRAFT"
    item.remarks = f"退回原因：{data.reason}" + (f"；{item.remarks}" if item.remarks else "")
    item.updated_by = user.id
    item.version += 1
    audit(
        db,
        user,
        request,
        "exam.result.return",
        "exam_attempt",
        item.id,
        {"reason": data.reason},
    )
    db.commit()
    return attempt_view(db, item)


@router.post("/exam-attempts/{attempt_id}/revise-result")
def revise_result(
    attempt_id: str, data: RevisionInput, request: Request, db: Db, user: ResultRevise
) -> dict[str, Any]:
    item = get_attempt_for_user(db, user, attempt_id)
    if item.version != data.version or item.result_confirm_status != "OFFICIALLY_CONFIRMED":
        raise error(409, "EXAM_RESULT_REVISION_CONFLICT", "只能修订官方确认结果。")
    subject = db.get(CourseSubjectVersion, item.course_subject_version_id)
    if data.score_value_scaled is not None and data.score_value_scaled > 100 * data.score_scale:
        raise error(422, "EXAM_SCORE_INVALID", "分数超出范围。")
    if data.score_value_scaled is not None and data.result_status in {"PASSED", "FAILED"}:
        suggested = (
            "PASSED"
            if subject.passing_score is not None
            and data.score_value_scaled >= subject.passing_score * data.score_scale
            else "FAILED"
        )
        if data.result_status != suggested:
            raise error(422, "EXAM_RESULT_SCORE_MISMATCH", "分数与结果不一致。")
    before = snapshot(item)
    item.score_value_scaled = data.score_value_scaled
    item.score_scale = data.score_scale
    item.result_status = data.result_status
    item.attendance_status = data.attendance_status
    item.remarks = data.remarks
    item.current_revision_no += 1
    item.result_confirm_status = "OFFICIALLY_CONFIRMED"
    item.confirmed_at = utc_now()
    item.confirmed_by = user.id
    item.version += 1
    after = snapshot(item)
    db.add(
        ExamResultRevision(
            exam_attempt_id=item.id,
            revision_no=item.current_revision_no,
            previous_values_json=json.dumps(before),
            new_values_json=json.dumps(after),
            reason=data.reason,
            request_id=request.state.correlation_id,
            revised_by=user.id,
        )
    )
    enrollment = db.get(CourseEnrollment, item.course_enrollment_id)
    status = recompute(db, enrollment)
    certificate = db.scalar(
        select(CertificateCase).where(CertificateCase.course_enrollment_id == enrollment.id)
    )
    if certificate and status != "PASSED":
        before_certificate = certificate.certificate_status
        certificate.eligibility_status = "INVALIDATED_REVIEW_REQUIRED"
        certificate.certificate_status = "EXCEPTION"
        certificate.exception_reason = "官方成绩修订后不再满足证书资格，需要人工复核。"
        certificate.updated_by = user.id
        certificate.version += 1
        enrollment.certificate_status = "EXCEPTION"
        db.add(
            CertificateDeliveryEvent(
                certificate_case_id=certificate.id,
                event_type="ELIGIBILITY_INVALIDATED",
                from_status=before_certificate,
                to_status="EXCEPTION",
                operator_id=user.id,
                remarks=certificate.exception_reason,
            )
        )
    audit(db, user, request, "exam.result.revise", "exam_attempt", item.id, {"reason": data.reason})
    db.commit()
    return attempt_view(db, item)


@router.get("/exam-attempts/{attempt_id}/revisions")
def revisions(attempt_id: str, db: Db, _user: ResultRead) -> dict[str, Any]:
    get_attempt_for_user(db, _user, attempt_id)
    return {
        "items": [
            {
                "id": x.id,
                "revision_no": x.revision_no,
                "reason": x.reason,
                "created_at": x.created_at,
            }
            for x in db.scalars(
                select(ExamResultRevision)
                .where(ExamResultRevision.exam_attempt_id == attempt_id)
                .order_by(ExamResultRevision.revision_no)
            ).all()
        ]
    }


EXAM_IMPORT_HEADERS = [
    "考试批次编码",
    "报名编号",
    "姓名",
    "科目编码",
    "准考证号",
    "考试日期",
    "出席状态",
    "分数",
    "结果",
    "结果发布日期",
    "备注",
]


def parse_import_score(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        score = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("分数格式无效。") from exc
    if score < 0 or score > 100 or score.as_tuple().exponent < -2:
        raise ValueError("分数必须为 0 至 100，最多两位小数。")
    return int(score * 100)


def import_batch_for_user(db: Session, user: User, batch_id: str) -> ExamResultImportBatch:
    batch = db.scalar(
        select(ExamResultImportBatch).where(
            ExamResultImportBatch.id == batch_id,
            ExamResultImportBatch.organization_id == user.organization_id,
            ExamResultImportBatch.created_by == user.id,
        )
    )
    if not batch:
        raise error(404, "EXAM_IMPORT_BATCH_NOT_FOUND", "未找到成绩导入预检批次。")
    return batch


def exam_import_view(db: Session, batch: ExamResultImportBatch) -> dict[str, Any]:
    lines = db.scalars(
        select(ExamResultImportLine)
        .where(ExamResultImportLine.batch_id == batch.id)
        .order_by(ExamResultImportLine.row_number)
    ).all()
    return {
        "id": batch.id,
        "status": batch.status,
        "expires_at": batch.expires_at,
        "summary": json.loads(batch.summary_json),
        "items": [
            {
                "row_number": line.row_number,
                "status": line.status,
                "exam_attempt_id": line.exam_attempt_id,
                "result": json.loads(line.result_json),
            }
            for line in lines
        ],
    }


@router.get("/imports/exam-results/template")
def exam_import_template(_user: ResultImport) -> StreamingResponse:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "考试成绩"
    worksheet.append(EXAM_IMPORT_HEADERS)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=exam-result-import-template.xlsx"},
    )


@router.post("/imports/exam-results/preview")
async def exam_import_preview(
    request: Request,
    db: Db,
    user: ResultImport,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise error(422, "INVALID_IMPORT_FILE", "仅支持 xlsx 文件。")
    content = await file.read()
    settings = get_settings()
    if len(content) > settings.import_max_bytes:
        raise error(422, "IMPORT_FILE_TOO_LARGE", "导入文件超过大小限制。")
    try:
        worksheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
        headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
    except Exception as exc:
        raise error(422, "INVALID_IMPORT_FILE", "无法读取 Excel 文件。") from exc
    if headers != EXAM_IMPORT_HEADERS:
        raise error(422, "IMPORT_TEMPLATE_INVALID", "模板列不匹配，请下载最新模板。")

    seen: set[tuple[str, str, str]] = set()
    payloads: list[dict[str, Any]] = []
    line_values: list[tuple[int, str, str | None, dict[str, Any]]] = []
    counts: dict[str, int] = {}
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if row_number > settings.import_max_rows + 1:
            raise error(422, "IMPORT_TOO_MANY_ROWS", "导入行数超过限制。")
        values = dict(zip(EXAM_IMPORT_HEADERS, row, strict=True))
        batch_code = str(values["考试批次编码"] or "").strip()
        registration_no = str(values["报名编号"] or "").strip()
        subject_code = str(values["科目编码"] or "").strip()
        key = (batch_code, registration_no, subject_code)
        status = "READY"
        attempt: ExamAttempt | None = None
        result: dict[str, Any] = {"message": "可导入并进入待复核状态。"}
        payload: dict[str, Any] = {}
        try:
            if not all(key):
                raise ValueError("考试批次编码、报名编号和科目编码不能为空。")
            if key in seen:
                status = "DUPLICATE"
                result = {"message": "文件内存在重复行。"}
            else:
                seen.add(key)
                record = db.execute(
                    select(
                        ExamAttempt,
                        ExamRegistration,
                        CourseSubjectVersion,
                        CourseEnrollment,
                        Student,
                    )
                    .join(
                        ExamRegistrationSubject,
                        ExamAttempt.exam_registration_subject_id == ExamRegistrationSubject.id,
                    )
                    .join(
                        ExamRegistration,
                        ExamRegistrationSubject.exam_registration_id == ExamRegistration.id,
                    )
                    .join(ExamBatch, ExamRegistration.exam_batch_id == ExamBatch.id)
                    .join(
                        CourseSubjectVersion,
                        ExamAttempt.course_subject_version_id == CourseSubjectVersion.id,
                    )
                    .join(
                        CourseEnrollment,
                        ExamAttempt.course_enrollment_id == CourseEnrollment.id,
                    )
                    .join(Student, CourseEnrollment.student_id == Student.id)
                    .where(
                        ExamBatch.batch_code == batch_code,
                        ExamRegistration.registration_no == registration_no,
                        CourseSubjectVersion.subject_code == subject_code,
                    )
                ).first()
                if not record:
                    raise ValueError("未找到匹配的考试报名科目。")
                attempt, registration, subject, _enrollment, student = record
                score = parse_import_score(values["分数"])
                raw_result = str(values["结果"] or "").strip().upper()
                result_map = {"合格": "PASSED", "不合格": "FAILED", "缺考": "ABSENT"}
                result_status = result_map.get(raw_result, raw_result)
                if result_status not in {"PASSED", "FAILED", "ABSENT", "INVALID"}:
                    raise ValueError("结果必须是 PASSED、FAILED、ABSENT 或 INVALID。")
                attendance = str(values["出席状态"] or "PRESENT").strip().upper()
                attendance_map = {"出席": "PRESENT", "缺考": "ABSENT"}
                attendance = attendance_map.get(attendance, attendance)
                if attendance not in {
                    "PRESENT",
                    "ABSENT",
                    "EXCUSED_ABSENCE",
                    "DISQUALIFIED",
                    "UNKNOWN",
                }:
                    raise ValueError("出席状态无效。")
                suggested = (
                    "PASSED"
                    if score is not None
                    and subject.passing_score is not None
                    and score >= subject.passing_score * 100
                    else "FAILED"
                )
                if score is not None and result_status in {"PASSED", "FAILED"}:
                    if suggested != result_status:
                        raise ValueError("分数与结果不一致。")
                if attempt.result_confirm_status == "OFFICIALLY_CONFIRMED":
                    status = "CONFLICT"
                    result = {"message": "已有官方确认成绩，普通导入不会覆盖。"}
                elif str(values["姓名"] or "").strip() not in {"", student.full_name}:
                    status = "WARNING"
                    result = {"message": "姓名与报名学员不一致，请人工核对。"}
                payload = {
                    "exam_attempt_id": attempt.id,
                    "score_value_scaled": score,
                    "score_scale": 100,
                    "result_status": result_status,
                    "attendance_status": attendance,
                    "remarks": str(values["备注"] or "")[:500],
                }
        except ValueError as exc:
            status = "INVALID"
            result = {"message": str(exc)}
        payloads.append(payload)
        line_values.append((row_number, status, attempt.id if attempt else None, result))
        counts[status] = counts.get(status, 0) + 1

    batch = ExamResultImportBatch(
        organization_id=user.organization_id,
        created_by=user.id,
        expires_at=utc_now() + timedelta(hours=2),
        preview_payload_ciphertext=encrypt(json.dumps(payloads, ensure_ascii=False)),
        summary_json=json.dumps({"total": len(line_values), "counts": counts}),
    )
    db.add(batch)
    db.flush()
    for row_number, status, attempt_id, result in line_values:
        db.add(
            ExamResultImportLine(
                batch_id=batch.id,
                row_number=row_number,
                status=status,
                exam_attempt_id=attempt_id,
                result_json=json.dumps(result, ensure_ascii=False),
            )
        )
    audit(
        db,
        user,
        request,
        "exam.result.import.preview",
        "exam_result_import_batch",
        batch.id,
        {"rows": len(line_values), "counts": counts},
    )
    db.commit()
    return exam_import_view(db, batch)


@router.get("/imports/exam-results/{batch_id}")
def get_exam_import(batch_id: str, db: Db, user: ResultImport) -> dict[str, Any]:
    return exam_import_view(db, import_batch_for_user(db, user, batch_id))


@router.post("/imports/exam-results/{batch_id}/confirm")
def confirm_exam_import(
    batch_id: str,
    data: ImportConfirmInput,
    request: Request,
    db: Db,
    user: ResultImport,
) -> dict[str, Any]:
    batch = import_batch_for_user(db, user, batch_id)
    if batch.confirm_idempotency_key == data.idempotency_key:
        return exam_import_view(db, batch)
    if batch.status != "PREVIEWED":
        raise error(409, "EXAM_IMPORT_NOT_CONFIRMABLE", "预检批次不能再次确认。")
    if batch.expires_at < utc_now():
        raise error(409, "EXAM_IMPORT_EXPIRED", "预检批次已过期，请重新上传。")
    payloads = json.loads(decrypt(batch.preview_payload_ciphertext))
    lines = db.scalars(
        select(ExamResultImportLine)
        .where(ExamResultImportLine.batch_id == batch.id)
        .order_by(ExamResultImportLine.row_number)
    ).all()
    for payload, line in zip(payloads, lines, strict=True):
        if line.status != "READY":
            if line.status == "WARNING":
                line.status = "SKIPPED"
                line.result_json = json.dumps(
                    {"message": "警告行未自动导入，请核对后重新提交。"}, ensure_ascii=False
                )
            continue
        try:
            with db.begin_nested():
                attempt = db.get(ExamAttempt, payload["exam_attempt_id"])
                if not attempt or attempt.result_confirm_status == "OFFICIALLY_CONFIRMED":
                    raise ValueError("成绩已变更或已经官方确认。")
                attempt.score_value_scaled = payload["score_value_scaled"]
                attempt.score_scale = payload["score_scale"]
                attempt.result_status = payload["result_status"]
                attempt.attendance_status = payload["attendance_status"]
                attempt.remarks = payload["remarks"]
                attempt.result_source = "IMPORT"
                attempt.result_confirm_status = "SUBMITTED"
                attempt.updated_by = user.id
                attempt.version += 1
                line.status = "SUCCESS"
                line.result_json = json.dumps(
                    {"message": "已导入，等待官方确认。"}, ensure_ascii=False
                )
        except (ValueError, IntegrityError) as exc:
            line.status = "FAILED"
            line.result_json = json.dumps({"message": str(exc)}, ensure_ascii=False)
    batch.status = "CONFIRMED"
    batch.confirm_idempotency_key = data.idempotency_key
    audit(
        db,
        user,
        request,
        "exam.result.import.confirm",
        "exam_result_import_batch",
        batch.id,
    )
    db.commit()
    return exam_import_view(db, batch)


@router.get("/imports/exam-results/{batch_id}/result")
def exam_import_result(batch_id: str, db: Db, user: ResultImport) -> dict[str, Any]:
    return exam_import_view(db, import_batch_for_user(db, user, batch_id))


@router.get("/exam-fee-assessments")
def fees(db: Db, _user: FeeRead) -> dict[str, Any]:
    return {
        "items": [
            {
                "id": x.id,
                "assessment_no": x.assessment_no,
                "course_enrollment_id": x.course_enrollment_id,
                "fee_type": x.fee_type,
                "responsibility": x.responsibility,
                "amount_cent": x.amount_cent,
                "assessment_status": x.assessment_status,
                "version": x.version,
            }
            for x in db.scalars(
                select(ExamFeeAssessment)
                .join(CourseEnrollment)
                .where(CourseEnrollment.organization_id == _user.organization_id)
                .order_by(ExamFeeAssessment.created_at.desc())
            ).all()
        ]
    }


@router.get("/exam-fee-assessments/{assessment_id}")
def fee_detail(assessment_id: str, db: Db, _user: FeeRead) -> dict[str, Any]:
    item = db.scalar(
        select(ExamFeeAssessment)
        .join(CourseEnrollment)
        .where(
            ExamFeeAssessment.id == assessment_id,
            CourseEnrollment.organization_id == _user.organization_id,
        )
    )
    if not item:
        raise error(404, "FEE_ASSESSMENT_NOT_FOUND", "未找到费用认定。")
    return {
        "id": item.id,
        "assessment_no": item.assessment_no,
        "course_enrollment_id": item.course_enrollment_id,
        "exam_attempt_id": item.exam_attempt_id,
        "course_subject_version_id": item.course_subject_version_id,
        "fee_policy_version_id": item.fee_policy_version_id,
        "fee_type": item.fee_type,
        "responsibility": item.responsibility,
        "amount_cent": item.amount_cent,
        "assessment_status": item.assessment_status,
        "policy_snapshot_description": item.policy_snapshot_description,
        "finance_reference": item.finance_reference,
        "version": item.version,
    }


def fee_action(
    assessment_id: str, data: FeeAction, request: Request, db: Session, user: User, target: str
) -> dict[str, Any]:
    item = db.scalar(
        select(ExamFeeAssessment)
        .join(CourseEnrollment)
        .where(
            ExamFeeAssessment.id == assessment_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if (
        not item
        or item.version != data.version
        or item.assessment_status not in {"DRAFT", "CONFIRMED"}
    ):
        raise error(409, "FEE_ASSESSMENT_CONFLICT", "费用认定当前不能操作。")
    if target == "WAIVED" and not data.reason:
        raise error(422, "REASON_REQUIRED", "减免必须填写原因。")
    item.assessment_status = target
    item.version += 1
    if target == "CONFIRMED":
        item.confirmed_at = utc_now()
        item.confirmed_by = user.id
    if target == "WAIVED":
        item.waiver_reason = data.reason
    if target == "CANCELLED":
        item.cancellation_reason = data.reason
    audit(db, user, request, f"exam.fee.{target.lower()}", "exam_fee_assessment", item.id)
    db.commit()
    return {"id": item.id, "assessment_status": item.assessment_status, "version": item.version}


@router.post("/exam-fee-assessments/{assessment_id}/confirm")
def confirm_fee(assessment_id: str, data: FeeAction, request: Request, db: Db, user: FeeConfirm):
    return fee_action(assessment_id, data, request, db, user, "CONFIRMED")


@router.post("/exam-fee-assessments/{assessment_id}/waive")
def waive_fee(assessment_id: str, data: FeeAction, request: Request, db: Db, user: FeeWaive):
    return fee_action(assessment_id, data, request, db, user, "WAIVED")


@router.post("/exam-fee-assessments/{assessment_id}/cancel")
def cancel_fee(assessment_id: str, data: FeeAction, request: Request, db: Db, user: FeeConfirm):
    return fee_action(assessment_id, data, request, db, user, "CANCELLED")


def eligibility(db: Session, enrollment: CourseEnrollment) -> dict[str, Any]:
    subjects = db.scalars(
        select(CourseSubjectVersion).where(
            CourseSubjectVersion.exam_scheme_version_id == enrollment.exam_scheme_version_id,
            CourseSubjectVersion.enabled.is_(True),
        )
    ).all()
    passed = set(
        db.scalars(
            select(ExamAttempt.course_subject_version_id).where(
                ExamAttempt.course_enrollment_id == enrollment.id,
                ExamAttempt.result_status == "PASSED",
                ExamAttempt.result_confirm_status == "OFFICIALLY_CONFIRMED",
            )
        ).all()
    )
    missing = [
        {"id": x.id, "code": x.subject_code, "name": x.subject_name}
        for x in subjects
        if x.id not in passed
    ]
    return {
        "eligible": bool(subjects) and not missing,
        "evaluated_at": utc_now(),
        "exam_scheme_version_id": enrollment.exam_scheme_version_id,
        "required_subject_count": len(subjects),
        "passed_subject_count": len(subjects) - len(missing),
        "missing_subjects": missing,
        "blocking_reasons": [] if not missing else ["REQUIRED_SUBJECT_NOT_PASSED"],
    }


@router.get("/enrollments/{enrollment_id}/certificate-eligibility")
def certificate_eligibility(enrollment_id: str, db: Db, _user: CertificateRead) -> dict[str, Any]:
    item = db.get(CourseEnrollment, enrollment_id)
    if not item or item.organization_id != _user.organization_id:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名。")
    return eligibility(db, item)


@router.post("/enrollments/{enrollment_id}/certificate-case", status_code=201)
def create_certificate(
    enrollment_id: str, request: Request, db: Db, user: CertificateManage
) -> dict[str, Any]:
    enrollment = db.get(CourseEnrollment, enrollment_id)
    if not enrollment or enrollment.organization_id != user.organization_id:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名。")
    old = db.scalar(
        select(CertificateCase).where(CertificateCase.course_enrollment_id == enrollment_id)
    )
    if old:
        return {"id": old.id, "certificate_status": old.certificate_status, "idempotent": True}
    if not eligibility(db, enrollment)["eligible"]:
        raise error(409, "CERTIFICATE_NOT_ELIGIBLE", "尚未满足证书资格。")
    item = CertificateCase(
        certificate_case_no=new_code("CERT"),
        course_enrollment_id=enrollment.id,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(item)
    db.flush()
    enrollment.certificate_status = "ELIGIBLE"
    db.add(
        CertificateDeliveryEvent(
            certificate_case_id=item.id,
            event_type="ELIGIBILITY_CONFIRMED",
            from_status="ELIGIBLE",
            to_status="ELIGIBLE",
            operator_id=user.id,
            remarks="全部必考科目已官方确认通过。",
        )
    )
    audit(db, user, request, "certificate.create", "certificate_case", item.id)
    db.commit()
    return {"id": item.id, "certificate_status": item.certificate_status, "idempotent": False}


@router.get("/certificates")
def certificates(db: Db, _user: CertificateRead) -> dict[str, Any]:
    return {
        "items": [
            {
                "id": x.id,
                "certificate_case_no": x.certificate_case_no,
                "course_enrollment_id": x.course_enrollment_id,
                "eligibility_status": x.eligibility_status,
                "certificate_status": x.certificate_status,
                "certificate_no_masked": f"****{x.certificate_no_last4}"
                if x.certificate_no_last4
                else None,
                "version": x.version,
            }
            for x in db.scalars(
                select(CertificateCase)
                .join(CourseEnrollment)
                .where(CourseEnrollment.organization_id == _user.organization_id)
                .order_by(CertificateCase.created_at.desc())
            ).all()
        ]
    }


def certificate_for_user(db: Session, user: User, certificate_id: str) -> CertificateCase:
    item = db.scalar(
        select(CertificateCase)
        .join(CourseEnrollment)
        .where(
            CertificateCase.id == certificate_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not item:
        raise error(404, "CERTIFICATE_NOT_FOUND", "未找到证书案例。")
    return item


@router.get("/certificates/{certificate_id}")
def certificate_detail(certificate_id: str, db: Db, _user: CertificateRead) -> dict[str, Any]:
    item = certificate_for_user(db, _user, certificate_id)
    attachments = db.execute(
        select(ExamDomainAttachment, FileObject)
        .join(FileObject, ExamDomainAttachment.file_object_id == FileObject.id)
        .where(
            ExamDomainAttachment.owner_type == "CERTIFICATE_CASE",
            ExamDomainAttachment.owner_id == item.id,
        )
    ).all()
    return {
        "id": item.id,
        "certificate_case_no": item.certificate_case_no,
        "course_enrollment_id": item.course_enrollment_id,
        "eligibility_status": item.eligibility_status,
        "certificate_status": item.certificate_status,
        "certificate_type_code": item.certificate_type_code,
        "certificate_type_name": item.certificate_type_name_snapshot,
        "issuing_authority": item.issuing_authority,
        "certificate_no_masked": (
            f"****{item.certificate_no_last4}" if item.certificate_no_last4 else None
        ),
        "application_submitted_at": item.application_submitted_at,
        "issued_at": item.issued_at,
        "school_received_at": item.school_received_at,
        "ready_at": item.ready_at,
        "delivered_at": item.delivered_at,
        "delivery_method": item.delivery_method,
        "recipient_name": item.recipient_name,
        "courier_company": item.courier_company,
        "exception_reason": item.exception_reason,
        "version": item.version,
        "attachments": [
            {
                "id": link.id,
                "file_object_id": file.id,
                "attachment_type": link.attachment_type,
                "original_filename": file.original_filename,
            }
            for link, file in attachments
        ],
    }


@router.post("/certificates/{certificate_id}/reveal")
def reveal_certificate_numbers(
    certificate_id: str,
    data: RevealInput,
    request: Request,
    db: Db,
    user: CertificateSensitive,
) -> dict[str, Any]:
    item = certificate_for_user(db, user, certificate_id)
    audit(
        db,
        user,
        request,
        "certificate.sensitive.reveal",
        "certificate_case",
        item.id,
        {"reason": data.reason},
    )
    db.commit()
    return {
        "certificate_no": (
            decrypt(item.certificate_no_ciphertext) if item.certificate_no_ciphertext else None
        ),
        "tracking_no": decrypt(item.tracking_no_ciphertext)
        if item.tracking_no_ciphertext
        else None,
    }


def add_domain_attachment(
    owner_type: str,
    owner_id: str,
    attachment_type: str,
    upload: UploadFile,
    request: Request,
    db: Session,
    user: User,
) -> dict[str, Any]:
    filename = clean_filename(upload.filename)
    suffix = extension(filename)
    detected = detect(upload, suffix)
    settings = get_settings()
    limit = (
        settings.upload_max_pdf_bytes
        if detected == "application/pdf"
        else settings.upload_max_image_bytes
    )
    if upload.size is not None and upload.size > limit:
        raise error(413, "FILE_TOO_LARGE", "附件超过大小限制。")
    key = f"exam-attachments/{utc_now():%Y/%m}/{uuid.uuid4()}.{suffix}"
    storage = get_storage()
    try:
        stat = storage.put_stream(key, upload.file)
    except StorageError as exc:
        raise error(503, exc.code, "附件存储失败。") from exc
    if stat.size_bytes > limit:
        storage.delete(key)
        raise error(413, "FILE_TOO_LARGE", "附件超过大小限制。")
    file = FileObject(
        original_filename=filename,
        safe_extension=suffix,
        declared_content_type=upload.content_type,
        detected_content_type=detected,
        size_bytes=stat.size_bytes,
        sha256=stat.etag or "",
        file_status="ACTIVE",
        sensitivity_level="SENSITIVE",
        activated_at=utc_now(),
        created_by=user.id,
    )
    db.add(file)
    db.flush()
    db.add(
        FileReplica(
            file_object_id=file.id,
            storage_provider="LOCAL",
            object_key=key,
            replica_status="VERIFIED",
            is_primary=True,
            size_bytes=stat.size_bytes,
            sha256=file.sha256,
            verified_at=utc_now(),
            created_by=user.id,
        )
    )
    item = ExamDomainAttachment(
        owner_type=owner_type,
        owner_id=owner_id,
        file_object_id=file.id,
        attachment_type=attachment_type,
        created_by=user.id,
    )
    db.add(item)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise error(409, "ATTACHMENT_ALREADY_LINKED", "该附件已关联。")
    audit(db, user, request, "exam.attachment.link", owner_type.lower(), owner_id)
    db.commit()
    return {"id": item.id, "file_object_id": file.id}


@router.post("/exam-attempts/{attempt_id}/attachments", status_code=201)
def add_attempt_attachment(
    attempt_id: str,
    request: Request,
    db: Db,
    user: ResultEdit,
    attachment_type: str = Form(..., min_length=1, max_length=40),
    file: UploadFile = File(...),
) -> dict[str, Any]:
    get_attempt_for_user(db, user, attempt_id)
    return add_domain_attachment(
        "EXAM_ATTEMPT", attempt_id, attachment_type, file, request, db, user
    )


@router.post("/certificates/{certificate_id}/attachments", status_code=201)
def add_certificate_attachment(
    certificate_id: str,
    request: Request,
    db: Db,
    user: CertificateManage,
    attachment_type: str = Form(..., min_length=1, max_length=40),
    file: UploadFile = File(...),
) -> dict[str, Any]:
    certificate_for_user(db, user, certificate_id)
    return add_domain_attachment(
        "CERTIFICATE_CASE", certificate_id, attachment_type, file, request, db, user
    )


def download_domain_attachment(
    owner_type: str,
    owner_id: str,
    attachment_id: str,
    db: Session,
) -> StreamingResponse:
    row = db.execute(
        select(ExamDomainAttachment, FileObject, FileReplica)
        .join(FileObject, ExamDomainAttachment.file_object_id == FileObject.id)
        .join(
            FileReplica,
            FileReplica.file_object_id == FileObject.id,
        )
        .where(
            ExamDomainAttachment.id == attachment_id,
            ExamDomainAttachment.owner_type == owner_type,
            ExamDomainAttachment.owner_id == owner_id,
            FileReplica.is_primary.is_(True),
            FileReplica.deleted_at.is_(None),
        )
    ).first()
    if not row:
        raise error(404, "ATTACHMENT_NOT_FOUND", "未找到附件。")
    _link, file, replica = row
    try:
        stream = get_storage().open_stream(replica.object_key)
    except StorageError as exc:
        raise error(404, exc.code, "附件内容不可用。") from exc
    disposition = f"attachment; filename*=UTF-8''{quote(file.original_filename)}"
    return StreamingResponse(
        stream,
        media_type=file.detected_content_type,
        headers={"Content-Disposition": disposition},
    )


@router.get("/exam-attempts/{attempt_id}/attachments/{attachment_id}/download")
def download_attempt_attachment(
    attempt_id: str,
    attachment_id: str,
    request: Request,
    db: Db,
    user: ResultRead,
) -> StreamingResponse:
    get_attempt_for_user(db, user, attempt_id)
    response = download_domain_attachment("EXAM_ATTEMPT", attempt_id, attachment_id, db)
    audit(db, user, request, "exam.attachment.download", "exam_attempt", attempt_id)
    db.commit()
    return response


@router.get("/certificates/{certificate_id}/attachments/{attachment_id}/download")
def download_certificate_attachment(
    certificate_id: str,
    attachment_id: str,
    request: Request,
    db: Db,
    user: CertificateAttachmentRead,
) -> StreamingResponse:
    certificate_for_user(db, user, certificate_id)
    response = download_domain_attachment("CERTIFICATE_CASE", certificate_id, attachment_id, db)
    audit(db, user, request, "certificate.attachment.download", "certificate_case", certificate_id)
    db.commit()
    return response


def certificate_transition(
    certificate_id: str,
    data: CertificateAction,
    request: Request,
    db: Session,
    user: User,
    target: str,
) -> dict[str, Any]:
    item = certificate_for_user(db, user, certificate_id)
    allowed = {
        "APPLYING": {"ELIGIBLE"},
        "ISSUED": {"APPLYING"},
        "RECEIVED_BY_SCHOOL": {"ISSUED"},
        "READY_FOR_DELIVERY": {"RECEIVED_BY_SCHOOL"},
        "DELIVERED": {"READY_FOR_DELIVERY"},
        "EXCEPTION": {"ELIGIBLE", "APPLYING", "ISSUED", "RECEIVED_BY_SCHOOL", "READY_FOR_DELIVERY"},
    }
    if not item or item.version != data.version or item.certificate_status not in allowed[target]:
        raise error(409, "CERTIFICATE_STATUS_CONFLICT", "证书状态不能执行该操作。")
    before = item.certificate_status
    now = utc_now()
    item.certificate_status = target
    item.version += 1
    if target == "APPLYING":
        item.application_submitted_at = now
    if target == "ISSUED":
        if not data.certificate_no:
            raise error(422, "CERTIFICATE_NO_REQUIRED", "发证必须填写证书编号。")
        item.issued_at = now
        item.certificate_no_ciphertext = encrypt(data.certificate_no)
        item.certificate_no_last4 = data.certificate_no[-4:]
        item.issuing_authority = data.issuing_authority
    if target == "RECEIVED_BY_SCHOOL":
        item.school_received_at = now
    if target == "READY_FOR_DELIVERY":
        item.ready_at = now
    if target == "DELIVERED":
        if not data.delivery_method or not data.recipient_name:
            raise error(422, "DELIVERY_DETAILS_REQUIRED", "发放必须填写方式和领取人。")
        item.delivered_at = now
        item.delivery_method = data.delivery_method
        item.recipient_name = data.recipient_name
        item.courier_company = data.courier_company
        item.tracking_no_ciphertext = encrypt(data.tracking_no) if data.tracking_no else None
    if target == "EXCEPTION":
        if not data.remarks:
            raise error(422, "REASON_REQUIRED", "异常必须填写原因。")
        item.exception_reason = data.remarks
    item.updated_by = user.id
    enrollment = db.get(CourseEnrollment, item.course_enrollment_id)
    enrollment.certificate_status = target
    db.add(
        CertificateDeliveryEvent(
            certificate_case_id=item.id,
            event_type=target,
            from_status=before,
            to_status=target,
            operator_id=user.id,
            remarks=data.remarks,
        )
    )
    audit(db, user, request, f"certificate.{target.lower()}", "certificate_case", item.id)
    db.commit()
    return {"id": item.id, "certificate_status": item.certificate_status, "version": item.version}


@router.post("/certificates/{certificate_id}/start-application")
def start_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateManage
):
    return certificate_transition(certificate_id, data, request, db, user, "APPLYING")


@router.post("/certificates/{certificate_id}/mark-issued")
def issue_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateManage
):
    return certificate_transition(certificate_id, data, request, db, user, "ISSUED")


@router.post("/certificates/{certificate_id}/receive-by-school")
def receive_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateManage
):
    return certificate_transition(certificate_id, data, request, db, user, "RECEIVED_BY_SCHOOL")


@router.post("/certificates/{certificate_id}/ready-for-delivery")
def ready_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateManage
):
    return certificate_transition(certificate_id, data, request, db, user, "READY_FOR_DELIVERY")


@router.post("/certificates/{certificate_id}/deliver")
def deliver_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateDeliver
):
    return certificate_transition(certificate_id, data, request, db, user, "DELIVERED")


@router.post("/certificates/{certificate_id}/mark-exception")
def exception_certificate(
    certificate_id: str, data: CertificateAction, request: Request, db: Db, user: CertificateManage
):
    return certificate_transition(certificate_id, data, request, db, user, "EXCEPTION")


@router.get("/certificates/{certificate_id}/events")
def certificate_events(certificate_id: str, db: Db, _user: CertificateRead) -> dict[str, Any]:
    certificate_for_user(db, _user, certificate_id)
    return {
        "items": [
            {
                "id": x.id,
                "event_type": x.event_type,
                "from_status": x.from_status,
                "to_status": x.to_status,
                "occurred_at": x.occurred_at,
                "remarks": x.remarks,
            }
            for x in db.scalars(
                select(CertificateDeliveryEvent)
                .where(CertificateDeliveryEvent.certificate_case_id == certificate_id)
                .order_by(CertificateDeliveryEvent.occurred_at)
            ).all()
        ]
    }
