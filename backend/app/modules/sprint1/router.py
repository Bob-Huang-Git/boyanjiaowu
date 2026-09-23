# ruff: noqa: B008, E501, E701, E702
import json
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.models import (
    AuditLog,
    ClassCycle,
    ClassMembership,
    CourseCatalog,
    CourseEnrollment,
    CourseOfferingVersion,
    CourseSubjectVersion,
    CurriculumVersion,
    ExamSchemeVersion,
    FeePolicyVersion,
    RefundPolicyVersion,
    Student,
    StudentIdentityDocument,
    StudentImportBatch,
    StudentImportLine,
    User,
    utc_now,
)
from app.core.pii import (
    blind_index,
    decrypt,
    encrypt,
    mask_document,
    mask_phone,
    normalize_document,
    normalize_phone,
)
from app.modules.iam.router import Db, require_permission

router = APIRouter(prefix="/api", tags=["sprint-1"])
StudentRead = Annotated[User, Depends(require_permission("student.read"))]
StudentCreate = Annotated[User, Depends(require_permission("student.create"))]
StudentUpdate = Annotated[User, Depends(require_permission("student.update"))]
StudentReveal = Annotated[User, Depends(require_permission("student.sensitive.reveal"))]
StudentImport = Annotated[User, Depends(require_permission("student.import"))]
CourseRead = Annotated[User, Depends(require_permission("course.read"))]
CourseManage = Annotated[User, Depends(require_permission("course.manage"))]
CoursePublish = Annotated[User, Depends(require_permission("course.publish"))]
EnrollmentRead = Annotated[User, Depends(require_permission("enrollment.read"))]
EnrollmentCreate = Annotated[User, Depends(require_permission("enrollment.create"))]
EnrollmentCancel = Annotated[User, Depends(require_permission("enrollment.cancel"))]
ClassRead = Annotated[User, Depends(require_permission("class.read"))]
ClassManage = Annotated[User, Depends(require_permission("class.manage"))]
ClassMemberManage = Annotated[User, Depends(require_permission("class.member.manage"))]


class IdentityInput(BaseModel):
    document_type: str = "PRC_ID"
    document_number: str = Field(min_length=2, max_length=100)
    is_primary: bool = True


class StudentInput(BaseModel):
    full_name: str = Field(min_length=1, max_length=120)
    gender: str = "UNKNOWN"
    birth_date: date | None = None
    phone: str | None = Field(default=None, max_length=40)
    email: str | None = None
    education_level: str | None = None
    residence_region_code: str | None = None
    source_channel: str | None = None
    remarks: str | None = None
    identity_document: IdentityInput | None = None
    force_create_reason: str | None = None


class StudentPatch(BaseModel):
    full_name: str | None = Field(default=None, min_length=1, max_length=120)
    gender: str | None = None
    birth_date: date | None = None
    phone: str | None = Field(default=None, max_length=40)
    email: str | None = None
    education_level: str | None = None
    residence_region_code: str | None = None
    source_channel: str | None = None
    remarks: str | None = None
    version: int


class RevealInput(BaseModel):
    reason: str = Field(min_length=1, max_length=300)


class CourseInput(BaseModel):
    course_code: str = Field(min_length=1, max_length=40)
    course_name: str = Field(min_length=1, max_length=120)
    category_code: str | None = None
    description: str | None = None


class CoursePatch(BaseModel):
    course_name: str | None = Field(default=None, min_length=1, max_length=120)
    category_code: str | None = None
    description: str | None = None
    enabled: bool | None = None
    version: int


class ComponentInput(BaseModel):
    version_no: str
    name: str
    total_minutes: int | None = None
    theory_minutes: int | None = None
    practical_minutes: int | None = None
    scheme_type: str = "STANDARD"
    pass_rule_description: str | None = None
    tuition_amount_cent: int | None = None
    initial_exam_fee_cent: int | None = None
    policy_description: str | None = None
    description: str | None = None
    subjects: list[dict[str, Any]] = []


class OfferingInput(BaseModel):
    version_no: str
    version_name: str
    effective_from: date | None = None
    effective_until: date | None = None
    curriculum: ComponentInput | None = None
    exam_scheme: ComponentInput | None = None
    fee_policy: ComponentInput | None = None
    refund_policy: ComponentInput | None = None


class OfferingPatch(BaseModel):
    version_name: str | None = Field(default=None, min_length=1, max_length=120)
    effective_from: date | None = None
    effective_until: date | None = None
    version: int


class EnrollmentInput(BaseModel):
    student_id: str
    course_offering_version_id: str
    enrollment_channel: str | None = None
    remarks: str | None = None
    idempotency_key: str = Field(min_length=8, max_length=100)


class CancelInput(BaseModel):
    reason: str = Field(min_length=1, max_length=500)


class ClassInput(BaseModel):
    class_code: str
    class_name: str
    course_offering_version_id: str
    planned_start_date: date | None = None
    planned_end_date: date | None = None
    capacity: int | None = Field(default=None, ge=1)
    class_status: str = "DRAFT"
    remarks: str | None = None
    version: int | None = None


class ClassPatch(BaseModel):
    class_name: str | None = Field(default=None, min_length=1, max_length=120)
    planned_start_date: date | None = None
    planned_end_date: date | None = None
    capacity: int | None = Field(default=None, ge=1)
    class_status: str | None = None
    remarks: str | None = None
    version: int


class MemberInput(BaseModel):
    course_enrollment_id: str
    entry_reason: str = "INITIAL"
    idempotency_key: str = Field(min_length=8, max_length=100)


class ConfirmInput(BaseModel):
    idempotency_key: str = Field(min_length=8, max_length=100)
    possible_duplicate_actions: dict[str, dict[str, str]] = {}


def error(status: int, code: str, message: str, **extra: Any) -> HTTPException:
    return HTTPException(status_code=status, detail={"code": code, "message": message, **extra})


def audit(
    db: Session,
    user: User,
    request: Request,
    action: str,
    subject_type: str,
    subject_id: str,
    detail: dict[str, Any] | None = None,
) -> None:
    db.add(
        AuditLog(
            organization_id=user.organization_id,
            actor_user_id=user.id,
            action=action,
            subject_type=subject_type,
            subject_id=subject_id,
            correlation_id=request.state.correlation_id,
            detail_json=json.dumps(detail or {}, ensure_ascii=False),
        )
    )


def code(prefix: str) -> str:
    from uuid import uuid4

    return f"{prefix}{uuid4().hex[:12].upper()}"


def document_summary(db: Session, student_id: str) -> str | None:
    document = db.scalar(
        select(StudentIdentityDocument).where(
            StudentIdentityDocument.student_id == student_id,
            StudentIdentityDocument.is_primary.is_(True),
        )
    )
    return (
        mask_document(document.document_number_last4, document.document_type) if document else None
    )


def student_view(db: Session, student: Student) -> dict[str, Any]:
    phone = mask_phone(decrypt(student.phone_ciphertext)) if student.phone_ciphertext else None
    enrollment_count = db.scalar(
        select(func.count())
        .select_from(CourseEnrollment)
        .where(CourseEnrollment.student_id == student.id)
    )
    return {
        "id": student.id,
        "student_no": student.student_no,
        "full_name": student.full_name,
        "gender": student.gender,
        "birth_date": student.birth_date,
        "phone_masked": phone,
        "document_summary": document_summary(db, student.id),
        "email": student.email,
        "education_level": student.education_level,
        "residence_region_code": student.residence_region_code,
        "source_channel": student.source_channel,
        "student_status": student.student_status,
        "remarks": student.remarks,
        "enrollment_count": enrollment_count or 0,
        "version": student.version,
        "updated_at": student.updated_at,
    }


def find_student(db: Session, user: User, student_id: str) -> Student:
    student = db.scalar(
        select(Student).where(
            Student.id == student_id, Student.organization_id == user.organization_id
        )
    )
    if not student:
        raise error(404, "STUDENT_NOT_FOUND", "未找到学员。")
    return student


def duplicate_result(db: Session, user: User, data: StudentInput) -> dict[str, Any]:
    confirmed: StudentIdentityDocument | None = None
    if data.identity_document:
        normalized = normalize_document(
            data.identity_document.document_type, data.identity_document.document_number
        )
        confirmed = db.scalar(
            select(StudentIdentityDocument).where(
                StudentIdentityDocument.document_type == data.identity_document.document_type,
                StudentIdentityDocument.document_number_blind_index == blind_index(normalized),
            )
        )
    if confirmed:
        existing = db.get(Student, confirmed.student_id)
        return {"status": "DUPLICATE_CONFIRMED", "student_id": existing.id if existing else None}
    possible = []
    statement = select(Student).where(
        Student.organization_id == user.organization_id, Student.full_name == data.full_name
    )
    if data.birth_date:
        statement = statement.where(Student.birth_date == data.birth_date)
    for candidate in db.scalars(statement.limit(10)):
        possible.append(
            {
                "student_id": candidate.id,
                "student_no": candidate.student_no,
                "full_name": candidate.full_name,
            }
        )
    return {"status": "POSSIBLE_DUPLICATE" if possible else "CLEAR", "candidates": possible}


def create_student(db: Session, user: User, request: Request, data: StudentInput) -> Student:
    duplicate = duplicate_result(db, user, data)
    if duplicate["status"] == "DUPLICATE_CONFIRMED":
        raise error(
            409, "DUPLICATE_CONFIRMED", "证件号已关联现有学员。", student_id=duplicate["student_id"]
        )
    if duplicate["status"] == "POSSIBLE_DUPLICATE" and not data.force_create_reason:
        raise error(
            409,
            "POSSIBLE_DUPLICATE",
            "存在疑似重复学员，请确认后填写原因。",
            candidates=duplicate["candidates"],
        )
    phone = normalize_phone(data.phone) if data.phone else None
    student = Student(
        organization_id=user.organization_id,
        student_no=code("STU"),
        full_name=data.full_name.strip(),
        gender=data.gender,
        birth_date=data.birth_date,
        phone_ciphertext=encrypt(phone) if phone else None,
        phone_blind_index=blind_index(phone) if phone else None,
        email=data.email,
        education_level=data.education_level,
        residence_region_code=data.residence_region_code,
        source_channel=data.source_channel,
        remarks=data.remarks,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(student)
    db.flush()
    if data.identity_document:
        normalized = normalize_document(
            data.identity_document.document_type, data.identity_document.document_number
        )
        db.add(
            StudentIdentityDocument(
                student_id=student.id,
                document_type=data.identity_document.document_type,
                document_number_ciphertext=encrypt(normalized),
                document_number_blind_index=blind_index(normalized),
                document_number_last4=normalized[-4:],
                is_primary=data.identity_document.is_primary,
                created_by=user.id,
                updated_by=user.id,
            )
        )
    detail = {"possible_duplicate_confirmed": bool(data.force_create_reason)}
    audit(db, user, request, "student.create", "student", student.id, detail)
    return student


@router.get("/students")
def list_students(
    db: Db,
    user: StudentRead,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    keyword: str | None = None,
    student_status: str | None = None,
    sort: str = "updated_at",
    order: str = "desc",
) -> dict[str, Any]:
    sort_columns = {
        "student_no": Student.student_no,
        "full_name": Student.full_name,
        "updated_at": Student.updated_at,
    }
    if sort not in sort_columns or order not in {"asc", "desc"}:
        raise error(422, "INVALID_SORT", "不支持的排序字段。")
    statement = select(Student).where(Student.organization_id == user.organization_id)
    if keyword:
        statement = statement.where(
            (Student.student_no == keyword) | Student.full_name.contains(keyword)
        )
    if student_status:
        statement = statement.where(Student.student_status == student_status)
    total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
    column = sort_columns[sort]
    statement = statement.order_by(column.desc() if order == "desc" else column.asc())
    students = db.scalars(statement.offset((page - 1) * page_size).limit(page_size)).all()
    return {
        "items": [student_view(db, student) for student in students],
        "page": page,
        "page_size": page_size,
        "total": total,
    }


@router.post("/students")
def post_student(
    data: StudentInput, request: Request, db: Db, user: StudentCreate
) -> dict[str, Any]:
    try:
        student = create_student(db, user, request, data)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "DUPLICATE_CONFIRMED", "证件号已关联现有学员。") from exc
    return student_view(db, student)


@router.post("/students/duplicate-check")
def check_duplicate(data: StudentInput, db: Db, user: StudentCreate) -> dict[str, Any]:
    return duplicate_result(db, user, data)


@router.get("/students/{student_id}")
def get_student(student_id: str, db: Db, user: StudentRead) -> dict[str, Any]:
    return student_view(db, find_student(db, user, student_id))


@router.patch("/students/{student_id}")
def patch_student(
    student_id: str, data: StudentPatch, request: Request, db: Db, user: StudentUpdate
) -> dict[str, Any]:
    student = find_student(db, user, student_id)
    if student.version != data.version:
        raise error(409, "VERSION_CONFLICT", "学员已被其他操作更新，请刷新后重试。")
    changed: list[str] = []
    for field, value in data.model_dump(exclude={"version"}, exclude_unset=True).items():
        if field == "phone":
            normalized = normalize_phone(value) if value else None
            student.phone_ciphertext = encrypt(normalized) if normalized else None
            student.phone_blind_index = blind_index(normalized) if normalized else None
            changed.append("phone_changed")
        else:
            setattr(student, field, value)
            changed.append(field)
    student.version += 1
    student.updated_by = user.id
    audit(db, user, request, "student.update", "student", student.id, {"changed": changed})
    db.commit()
    return student_view(db, student)


@router.post("/students/{student_id}/deactivate")
def deactivate_student(
    student_id: str, request: Request, db: Db, user: StudentUpdate
) -> dict[str, bool]:
    student = find_student(db, user, student_id)
    student.student_status = "INACTIVE"
    student.version += 1
    student.updated_by = user.id
    audit(db, user, request, "student.deactivate", "student", student.id)
    db.commit()
    return {"ok": True}


@router.post("/students/{student_id}/identity-documents/{document_id}/reveal")
def reveal_document(
    student_id: str,
    document_id: str,
    data: RevealInput,
    request: Request,
    db: Db,
    user: StudentReveal,
) -> dict[str, str]:
    find_student(db, user, student_id)
    document = db.scalar(
        select(StudentIdentityDocument).where(
            StudentIdentityDocument.id == document_id,
            StudentIdentityDocument.student_id == student_id,
        )
    )
    if not document:
        raise error(404, "DOCUMENT_NOT_FOUND", "未找到证件记录。")
    audit(
        db,
        user,
        request,
        "student.sensitive.reveal",
        "identity_document",
        document.id,
        {"reason": data.reason},
    )
    db.commit()
    return {
        "document_type": document.document_type,
        "document_number": decrypt(document.document_number_ciphertext),
    }


@router.get("/students/{student_id}/enrollments")
def student_enrollments(student_id: str, db: Db, user: EnrollmentRead) -> dict[str, Any]:
    find_student(db, user, student_id)
    rows = db.scalars(
        select(CourseEnrollment).where(CourseEnrollment.student_id == student_id)
    ).all()
    return {"items": [enrollment_view(db, row) for row in rows]}


@router.get("/students/{student_id}/class-memberships")
def student_memberships(student_id: str, db: Db, user: ClassRead) -> dict[str, Any]:
    find_student(db, user, student_id)
    rows = db.execute(
        select(ClassMembership, ClassCycle, CourseEnrollment)
        .join(CourseEnrollment, ClassMembership.course_enrollment_id == CourseEnrollment.id)
        .join(ClassCycle, ClassMembership.class_cycle_id == ClassCycle.id)
        .where(CourseEnrollment.student_id == student_id)
    ).all()
    return {
        "items": [
            {
                "id": membership.id,
                "class_id": cycle.id,
                "class_code": cycle.class_code,
                "class_name": cycle.class_name,
                "membership_status": membership.membership_status,
                "entered_at": membership.entered_at,
            }
            for membership, cycle, _enrollment in rows
        ]
    }


def course_view(course: CourseCatalog) -> dict[str, Any]:
    return {
        "id": course.id,
        "course_code": course.course_code,
        "course_name": course.course_name,
        "category_code": course.category_code,
        "description": course.description,
        "enabled": course.enabled,
        "version": course.version,
    }


def offering_view(db: Session, offering: CourseOfferingVersion) -> dict[str, Any]:
    count = (
        db.scalar(
            select(func.count())
            .select_from(CourseEnrollment)
            .where(CourseEnrollment.course_offering_version_id == offering.id)
        )
        or 0
    )
    return {
        "id": offering.id,
        "course_catalog_id": offering.course_catalog_id,
        "version_no": offering.version_no,
        "version_name": offering.version_name,
        "status": offering.status,
        "effective_from": offering.effective_from,
        "effective_until": offering.effective_until,
        "curriculum_version_id": offering.curriculum_version_id,
        "exam_scheme_version_id": offering.exam_scheme_version_id,
        "fee_policy_version_id": offering.fee_policy_version_id,
        "refund_policy_version_id": offering.refund_policy_version_id,
        "enrollment_count": count,
        "version": offering.version,
    }


def create_components(
    db: Session, data: OfferingInput
) -> tuple[str | None, str | None, str | None, str | None]:
    curriculum_id = exam_id = fee_id = refund_id = None
    if data.curriculum:
        if data.curriculum.total_minutes is None:
            raise error(422, "INVALID_CURRICULUM", "课程方案必须填写总分钟数。")
        curriculum = CurriculumVersion(
            version_no=data.curriculum.version_no,
            name=data.curriculum.name,
            total_minutes=data.curriculum.total_minutes,
            theory_minutes=data.curriculum.theory_minutes,
            practical_minutes=data.curriculum.practical_minutes,
            description=data.curriculum.description,
        )
        db.add(curriculum)
        db.flush()
        curriculum_id = curriculum.id
    if data.exam_scheme:
        exam = ExamSchemeVersion(
            version_no=data.exam_scheme.version_no,
            name=data.exam_scheme.name,
            scheme_type=data.exam_scheme.scheme_type,
            pass_rule_description=data.exam_scheme.pass_rule_description,
        )
        db.add(exam)
        db.flush()
        exam_id = exam.id
        for index, subject in enumerate(data.exam_scheme.subjects, start=1):
            db.add(
                CourseSubjectVersion(
                    exam_scheme_version_id=exam.id,
                    subject_code=str(subject["subject_code"]),
                    subject_name=str(subject["subject_name"]),
                    display_order=int(subject.get("display_order", index)),
                    passing_score=subject.get("passing_score"),
                )
            )
    if data.fee_policy:
        if data.fee_policy.tuition_amount_cent is None:
            raise error(422, "INVALID_FEE_POLICY", "收费规则必须填写培训费金额（分）。")
        fee = FeePolicyVersion(
            version_no=data.fee_policy.version_no,
            name=data.fee_policy.name,
            tuition_amount_cent=data.fee_policy.tuition_amount_cent,
            initial_exam_fee_cent=data.fee_policy.initial_exam_fee_cent,
            description=data.fee_policy.description,
        )
        db.add(fee)
        db.flush()
        fee_id = fee.id
    if data.refund_policy:
        refund = RefundPolicyVersion(
            version_no=data.refund_policy.version_no,
            name=data.refund_policy.name,
            policy_description=data.refund_policy.policy_description,
        )
        db.add(refund)
        db.flush()
        refund_id = refund.id
    return curriculum_id, exam_id, fee_id, refund_id


@router.get("/courses")
def list_courses(db: Db, user: CourseRead) -> dict[str, Any]:
    rows = db.scalars(
        select(CourseCatalog)
        .where(CourseCatalog.organization_id == user.organization_id)
        .order_by(CourseCatalog.course_code)
    ).all()
    return {"items": [course_view(row) for row in rows]}


@router.post("/courses")
def post_course(data: CourseInput, request: Request, db: Db, user: CourseManage) -> dict[str, Any]:
    course = CourseCatalog(
        organization_id=user.organization_id,
        course_code=data.course_code.upper(),
        course_name=data.course_name,
        category_code=data.category_code,
        description=data.description,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(course)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "COURSE_CODE_EXISTS", "课程编码已存在。") from exc
    audit(db, user, request, "course.create", "course", course.id)
    db.commit()
    return course_view(course)


@router.get("/courses/{course_id}")
def get_course(course_id: str, db: Db, user: CourseRead) -> dict[str, Any]:
    course = db.scalar(
        select(CourseCatalog).where(
            CourseCatalog.id == course_id, CourseCatalog.organization_id == user.organization_id
        )
    )
    if not course:
        raise error(404, "COURSE_NOT_FOUND", "未找到课程。")
    return course_view(course)


@router.patch("/courses/{course_id}")
def patch_course(
    course_id: str, data: CoursePatch, request: Request, db: Db, user: CourseManage
) -> dict[str, Any]:
    course = db.scalar(
        select(CourseCatalog).where(
            CourseCatalog.id == course_id, CourseCatalog.organization_id == user.organization_id
        )
    )
    if not course:
        raise error(404, "COURSE_NOT_FOUND", "未找到课程。")
    if course.version != data.version:
        raise error(409, "VERSION_CONFLICT", "课程已被其他操作更新，请刷新后重试。")
    for field, value in data.model_dump(exclude={"version"}, exclude_unset=True).items():
        setattr(course, field, value)
    course.version += 1
    course.updated_by = user.id
    audit(db, user, request, "course.update", "course", course.id)
    db.commit()
    return course_view(course)


@router.get("/courses/{course_id}/versions")
def list_offerings(course_id: str, db: Db, user: CourseRead) -> dict[str, Any]:
    get_course(course_id, db, user)
    rows = db.scalars(
        select(CourseOfferingVersion)
        .where(CourseOfferingVersion.course_catalog_id == course_id)
        .order_by(CourseOfferingVersion.version_no)
    ).all()
    return {"items": [offering_view(db, row) for row in rows]}


@router.post("/courses/{course_id}/versions")
def post_offering(
    course_id: str, data: OfferingInput, request: Request, db: Db, user: CourseManage
) -> dict[str, Any]:
    get_course(course_id, db, user)
    ids = create_components(db, data)
    offering = CourseOfferingVersion(
        course_catalog_id=course_id,
        version_no=data.version_no,
        version_name=data.version_name,
        effective_from=data.effective_from,
        effective_until=data.effective_until,
        curriculum_version_id=ids[0],
        exam_scheme_version_id=ids[1],
        fee_policy_version_id=ids[2],
        refund_policy_version_id=ids[3],
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(offering)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "COURSE_VERSION_EXISTS", "该课程版本号已存在。") from exc
    audit(db, user, request, "course.version.create", "course_version", offering.id)
    db.commit()
    return offering_view(db, offering)


@router.get("/course-versions/{version_id}")
def get_offering(version_id: str, db: Db, user: CourseRead) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, version_id)
    if not offering:
        raise error(404, "COURSE_VERSION_NOT_FOUND", "未找到课程版本。")
    get_course(offering.course_catalog_id, db, user)
    return offering_view(db, offering)


@router.patch("/course-versions/{version_id}")
def patch_offering(
    version_id: str, data: OfferingPatch, request: Request, db: Db, user: CourseManage
) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, version_id)
    if not offering:
        raise error(404, "COURSE_VERSION_NOT_FOUND", "未找到课程版本。")
    get_course(offering.course_catalog_id, db, user)
    if offering.status != "DRAFT":
        raise error(409, "VERSION_IMMUTABLE", "已发布或停用的版本不能直接修改，请创建新版本。")
    if offering.version != data.version:
        raise error(409, "VERSION_CONFLICT", "课程版本已被其他操作更新，请刷新后重试。")
    for field, value in data.model_dump(exclude={"version"}, exclude_unset=True).items():
        setattr(offering, field, value)
    offering.version += 1
    offering.updated_by = user.id
    audit(db, user, request, "course.version.update", "course_version", offering.id)
    db.commit()
    return offering_view(db, offering)


@router.post("/course-versions/{version_id}/publish")
def publish_offering(
    version_id: str, request: Request, db: Db, user: CoursePublish
) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, version_id)
    if not offering:
        raise error(404, "COURSE_VERSION_NOT_FOUND", "未找到课程版本。")
    get_course(offering.course_catalog_id, db, user)
    if offering.status != "DRAFT":
        raise error(409, "VERSION_IMMUTABLE", "只有草稿版本可以发布。")
    if not all(
        [
            offering.curriculum_version_id,
            offering.exam_scheme_version_id,
            offering.fee_policy_version_id,
            offering.refund_policy_version_id,
        ]
    ):
        raise error(422, "VERSION_INCOMPLETE", "课程版本缺少必需方案或规则。")
    subject_count = (
        db.scalar(
            select(func.count())
            .select_from(CourseSubjectVersion)
            .where(CourseSubjectVersion.exam_scheme_version_id == offering.exam_scheme_version_id)
        )
        or 0
    )
    if not subject_count:
        raise error(422, "VERSION_INCOMPLETE", "考试方案至少需要一个科目。")
    offering.status = "PUBLISHED"
    offering.published_at = utc_now()
    offering.published_by = user.id
    offering.updated_by = user.id
    offering.version += 1
    for component in [
        db.get(CurriculumVersion, offering.curriculum_version_id),
        db.get(ExamSchemeVersion, offering.exam_scheme_version_id),
        db.get(FeePolicyVersion, offering.fee_policy_version_id),
        db.get(RefundPolicyVersion, offering.refund_policy_version_id),
    ]:
        component.status = "PUBLISHED"
    audit(db, user, request, "course.version.publish", "course_version", offering.id)
    db.commit()
    return offering_view(db, offering)


@router.post("/course-versions/{version_id}/retire")
def retire_offering(
    version_id: str, request: Request, db: Db, user: CoursePublish
) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, version_id)
    if not offering:
        raise error(404, "COURSE_VERSION_NOT_FOUND", "未找到课程版本。")
    offering.status = "RETIRED"
    offering.version += 1
    offering.updated_by = user.id
    audit(db, user, request, "course.version.retire", "course_version", offering.id)
    db.commit()
    return offering_view(db, offering)


def enrollment_view(db: Session, enrollment: CourseEnrollment) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, enrollment.course_offering_version_id)
    course = db.get(CourseCatalog, offering.course_catalog_id) if offering else None
    return {
        "id": enrollment.id,
        "enrollment_no": enrollment.enrollment_no,
        "student_id": enrollment.student_id,
        "course_offering_version_id": enrollment.course_offering_version_id,
        "course_code": course.course_code if course else None,
        "course_name": course.course_name if course else None,
        "version_name": offering.version_name if offering else None,
        "enrolled_at": enrollment.enrolled_at,
        "lifecycle_status": enrollment.lifecycle_status,
        "learning_status": enrollment.learning_status,
        "exam_status": enrollment.exam_status,
        "financial_status": enrollment.financial_status,
        "certificate_status": enrollment.certificate_status,
        "funding_status": enrollment.funding_status,
        "version": enrollment.version,
    }


@router.get("/enrollments")
def list_enrollments(db: Db, user: EnrollmentRead, student_id: str | None = None) -> dict[str, Any]:
    statement = select(CourseEnrollment).where(
        CourseEnrollment.organization_id == user.organization_id
    )
    if student_id:
        statement = statement.where(CourseEnrollment.student_id == student_id)
    return {
        "items": [
            enrollment_view(db, row)
            for row in db.scalars(statement.order_by(CourseEnrollment.enrolled_at.desc())).all()
        ]
    }


@router.post("/enrollments")
def post_enrollment(
    data: EnrollmentInput, request: Request, db: Db, user: EnrollmentCreate
) -> dict[str, Any]:
    existing = db.scalar(
        select(CourseEnrollment).where(
            CourseEnrollment.organization_id == user.organization_id,
            CourseEnrollment.idempotency_key == data.idempotency_key,
        )
    )
    if existing:
        return enrollment_view(db, existing)
    find_student(db, user, data.student_id)
    offering = db.get(CourseOfferingVersion, data.course_offering_version_id)
    if not offering or offering.status != "PUBLISHED":
        raise error(422, "OFFERING_NOT_AVAILABLE", "只能报名已发布且可用的课程版本。")
    course = db.get(CourseCatalog, offering.course_catalog_id)
    if not course or course.organization_id != user.organization_id or not course.enabled:
        raise error(422, "OFFERING_NOT_AVAILABLE", "课程当前不可报名。")
    today = date.today()
    if (offering.effective_from and offering.effective_from > today) or (
        offering.effective_until and offering.effective_until < today
    ):
        raise error(422, "OFFERING_NOT_AVAILABLE", "课程版本当前不在有效期内。")
    enrollment = CourseEnrollment(
        organization_id=user.organization_id,
        enrollment_no=code("ENR"),
        student_id=data.student_id,
        course_offering_version_id=offering.id,
        curriculum_version_id=offering.curriculum_version_id,
        exam_scheme_version_id=offering.exam_scheme_version_id,
        fee_policy_version_id=offering.fee_policy_version_id,
        refund_policy_version_id=offering.refund_policy_version_id,
        enrollment_channel=data.enrollment_channel,
        remarks=data.remarks,
        idempotency_key=data.idempotency_key,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(enrollment)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "IDEMPOTENCY_CONFLICT", "重复报名请求正在处理或已完成。") from exc
    audit(db, user, request, "enrollment.create", "enrollment", enrollment.id)
    db.commit()
    return enrollment_view(db, enrollment)


@router.get("/enrollments/{enrollment_id}")
def get_enrollment(enrollment_id: str, db: Db, user: EnrollmentRead) -> dict[str, Any]:
    enrollment = db.scalar(
        select(CourseEnrollment).where(
            CourseEnrollment.id == enrollment_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not enrollment:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名记录。")
    return enrollment_view(db, enrollment)


@router.post("/enrollments/{enrollment_id}/cancel")
def cancel_enrollment(
    enrollment_id: str, data: CancelInput, request: Request, db: Db, user: EnrollmentCancel
) -> dict[str, Any]:
    enrollment = db.scalar(
        select(CourseEnrollment).where(
            CourseEnrollment.id == enrollment_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not enrollment:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名记录。")
    if enrollment.lifecycle_status == "CANCELLED":
        return enrollment_view(db, enrollment)
    enrollment.lifecycle_status = "CANCELLED"
    enrollment.cancelled_at = utc_now()
    enrollment.cancelled_by = user.id
    enrollment.cancel_reason = data.reason
    enrollment.version += 1
    audit(
        db, user, request, "enrollment.cancel", "enrollment", enrollment.id, {"reason": data.reason}
    )
    db.commit()
    return enrollment_view(db, enrollment)


def class_view(db: Session, cycle: ClassCycle) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, cycle.course_offering_version_id)
    course = db.get(CourseCatalog, offering.course_catalog_id) if offering else None
    active = (
        db.scalar(
            select(func.count())
            .select_from(ClassMembership)
            .where(
                ClassMembership.class_cycle_id == cycle.id,
                ClassMembership.membership_status == "ACTIVE",
            )
        )
        or 0
    )
    return {
        "id": cycle.id,
        "class_code": cycle.class_code,
        "class_name": cycle.class_name,
        "course_offering_version_id": cycle.course_offering_version_id,
        "course_name": course.course_name if course else None,
        "version_name": offering.version_name if offering else None,
        "planned_start_date": cycle.planned_start_date,
        "planned_end_date": cycle.planned_end_date,
        "capacity": cycle.capacity,
        "active_member_count": active,
        "class_status": cycle.class_status,
        "remarks": cycle.remarks,
        "version": cycle.version,
    }


@router.get("/classes")
def list_classes(db: Db, user: ClassRead) -> dict[str, Any]:
    rows = db.scalars(
        select(ClassCycle)
        .where(ClassCycle.organization_id == user.organization_id)
        .order_by(ClassCycle.updated_at.desc())
    ).all()
    return {"items": [class_view(db, row) for row in rows]}


@router.post("/classes")
def post_class(data: ClassInput, request: Request, db: Db, user: ClassManage) -> dict[str, Any]:
    offering = db.get(CourseOfferingVersion, data.course_offering_version_id)
    if (
        not offering
        or not db.get(CourseCatalog, offering.course_catalog_id)
        or db.get(CourseCatalog, offering.course_catalog_id).organization_id != user.organization_id
    ):
        raise error(422, "INVALID_COURSE_VERSION", "班级必须关联本机构课程版本。")
    cycle = ClassCycle(
        organization_id=user.organization_id,
        class_code=data.class_code.upper(),
        class_name=data.class_name,
        course_offering_version_id=data.course_offering_version_id,
        planned_start_date=data.planned_start_date,
        planned_end_date=data.planned_end_date,
        capacity=data.capacity,
        class_status=data.class_status,
        remarks=data.remarks,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(cycle)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "CLASS_CODE_EXISTS", "班级编码已存在。") from exc
    audit(db, user, request, "class.create", "class", cycle.id)
    db.commit()
    return class_view(db, cycle)


@router.get("/classes/{class_id}")
def get_class(class_id: str, db: Db, user: ClassRead) -> dict[str, Any]:
    cycle = db.scalar(
        select(ClassCycle).where(
            ClassCycle.id == class_id, ClassCycle.organization_id == user.organization_id
        )
    )
    if not cycle:
        raise error(404, "CLASS_NOT_FOUND", "未找到班级。")
    return class_view(db, cycle)


@router.patch("/classes/{class_id}")
def patch_class(
    class_id: str, data: ClassPatch, request: Request, db: Db, user: ClassManage
) -> dict[str, Any]:
    cycle = db.scalar(
        select(ClassCycle).where(
            ClassCycle.id == class_id, ClassCycle.organization_id == user.organization_id
        )
    )
    if not cycle:
        raise error(404, "CLASS_NOT_FOUND", "未找到班级。")
    if cycle.version != data.version:
        raise error(409, "VERSION_CONFLICT", "班级已被其他操作更新，请刷新后重试。")
    for field, value in data.model_dump(exclude={"version"}, exclude_unset=True).items():
        setattr(cycle, field, value)
    cycle.version += 1
    cycle.updated_by = user.id
    audit(db, user, request, "class.update", "class", cycle.id)
    db.commit()
    return class_view(db, cycle)


@router.get("/classes/{class_id}/members")
def list_members(class_id: str, db: Db, user: ClassRead) -> dict[str, Any]:
    get_class(class_id, db, user)
    rows = db.execute(
        select(ClassMembership, CourseEnrollment, Student)
        .join(CourseEnrollment, ClassMembership.course_enrollment_id == CourseEnrollment.id)
        .join(Student, CourseEnrollment.student_id == Student.id)
        .where(ClassMembership.class_cycle_id == class_id)
    ).all()
    return {
        "items": [
            {
                "id": member.id,
                "student_id": student.id,
                "student_no": student.student_no,
                "full_name": student.full_name,
                "phone_masked": mask_phone(decrypt(student.phone_ciphertext))
                if student.phone_ciphertext
                else None,
                "enrollment_no": enrollment.enrollment_no,
                "entered_at": member.entered_at,
                "membership_status": member.membership_status,
            }
            for member, enrollment, student in rows
        ]
    }


@router.post("/classes/{class_id}/members")
def add_member(
    class_id: str, data: MemberInput, request: Request, db: Db, user: ClassMemberManage
) -> dict[str, Any]:
    cycle = get_class(class_id, db, user)
    if cycle["class_status"] in {"COMPLETED", "CANCELLED"}:
        raise error(409, "CLASS_NOT_JOINABLE", "已完成或取消的班级不能新增成员。")
    enrollment = db.scalar(
        select(CourseEnrollment).where(
            CourseEnrollment.id == data.course_enrollment_id,
            CourseEnrollment.organization_id == user.organization_id,
        )
    )
    if not enrollment:
        raise error(404, "ENROLLMENT_NOT_FOUND", "未找到报名记录。")
    if enrollment.lifecycle_status != "ACTIVE":
        raise error(409, "ENROLLMENT_NOT_ACTIVE", "只有有效报名可以加入班级。")
    if enrollment.course_offering_version_id != cycle["course_offering_version_id"]:
        raise error(422, "CLASS_COURSE_MISMATCH", "报名课程版本与班级课程版本不一致。")
    existing = db.scalar(
        select(ClassMembership).where(
            ClassMembership.course_enrollment_id == enrollment.id,
            ClassMembership.membership_status == "ACTIVE",
        )
    )
    if existing:
        raise error(409, "ACTIVE_MEMBERSHIP_EXISTS", "该报名已在其他班级中。")
    if cycle["capacity"] is not None and cycle["active_member_count"] >= cycle["capacity"]:
        raise error(409, "CLASS_CAPACITY_REACHED", "班级已达到容量上限。")
    member = ClassMembership(
        class_cycle_id=class_id,
        course_enrollment_id=enrollment.id,
        entry_reason=data.entry_reason,
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(member)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise error(409, "ACTIVE_MEMBERSHIP_EXISTS", "该报名已在其他班级中。") from exc
    audit(
        db,
        user,
        request,
        "class.member.add",
        "class_membership",
        member.id,
        {"idempotency_key": data.idempotency_key},
    )
    db.commit()
    return {"id": member.id, "membership_status": member.membership_status}


@router.post("/classes/{class_id}/members/{membership_id}/remove")
def remove_member(
    class_id: str,
    membership_id: str,
    data: CancelInput,
    request: Request,
    db: Db,
    user: ClassMemberManage,
) -> dict[str, bool]:
    cycle = get_class(class_id, db, user)
    if cycle["class_status"] not in {"DRAFT", "OPEN"}:
        raise error(409, "CLASS_MEMBER_REMOVE_FORBIDDEN", "只有未开课班级允许撤出成员。")
    member = db.scalar(
        select(ClassMembership).where(
            ClassMembership.id == membership_id, ClassMembership.class_cycle_id == class_id
        )
    )
    if not member:
        raise error(404, "MEMBERSHIP_NOT_FOUND", "未找到班级经历。")
    if member.membership_status == "ACTIVE":
        member.membership_status = "EXITED"
        member.exited_at = utc_now()
        member.exit_reason = data.reason
        member.version += 1
        member.updated_by = user.id
        audit(
            db,
            user,
            request,
            "class.member.remove",
            "class_membership",
            member.id,
            {"reason": data.reason},
        )
        db.commit()
    return {"ok": True}


IMPORT_HEADERS = [
    "姓名",
    "证件类型",
    "证件号码",
    "手机号",
    "性别",
    "出生日期",
    "文化程度",
    "来源渠道",
    "课程编码",
    "课程版本",
    "报名日期",
    "班级编码",
    "备注",
]


def import_row(
    values: dict[str, Any], db: Session, user: User
) -> tuple[str, dict[str, Any], dict[str, Any]]:
    name = str(values.get("姓名") or "").strip()
    if not name:
        return "INVALID", {"field": "姓名", "message": "姓名为必填项。"}, {}
    birth_value = values.get("出生日期")
    try:
        birth = (
            birth_value.date()
            if isinstance(birth_value, datetime)
            else date.fromisoformat(str(birth_value))
            if birth_value
            else None
        )
    except ValueError:
        return "INVALID", {"field": "出生日期", "message": "出生日期格式无效。"}, {}
    data = StudentInput(
        full_name=name,
        gender=str(values.get("性别") or "UNKNOWN"),
        birth_date=birth,
        phone=str(values.get("手机号") or "").strip() or None,
        education_level=str(values.get("文化程度") or "").strip() or None,
        source_channel=str(values.get("来源渠道") or "").strip() or None,
        remarks=str(values.get("备注") or "").strip() or None,
        identity_document=IdentityInput(
            document_type=str(values.get("证件类型") or "PRC_ID"),
            document_number=str(values.get("证件号码") or "").strip(),
        )
        if values.get("证件号码")
        else None,
    )
    duplicate = duplicate_result(db, user, data)
    payload = data.model_dump(mode="json")
    if duplicate["status"] == "DUPLICATE_CONFIRMED":
        return (
            "DUPLICATE_CONFIRMED",
            {"message": "证件号已存在", "student_id": duplicate["student_id"]},
            payload,
        )
    if duplicate["status"] == "POSSIBLE_DUPLICATE":
        return (
            "POSSIBLE_DUPLICATE",
            {"message": "存在疑似重复，请逐行处理。", "candidates": duplicate["candidates"]},
            payload,
        )
    if any(values.get(column) for column in ("课程编码", "课程版本", "班级编码")):
        return (
            "READY_CREATE_STUDENT",
            {"message": "本阶段仅执行学员导入；课程、报名和班级列已预检但不会执行。"},
            payload,
        )
    return "READY_CREATE_STUDENT", {"message": "可创建学员。"}, payload


@router.get("/imports/students/template")
def import_template(_user: StudentImport) -> StreamingResponse:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学员导入"
    worksheet.append(IMPORT_HEADERS)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=student-import-template.xlsx"},
    )


@router.post("/imports/students/preview")
async def preview_import(
    request: Request, db: Db, user: StudentImport, file: UploadFile = File(...)
) -> dict[str, Any]:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise error(422, "INVALID_IMPORT_FILE", "仅支持 xlsx 文件。")
    content = await file.read()
    settings = get_settings()
    if len(content) > settings.import_max_bytes:
        raise error(422, "IMPORT_FILE_TOO_LARGE", "导入文件超过大小限制。")
    try:
        worksheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    except Exception as exc:
        raise error(422, "INVALID_IMPORT_FILE", "无法读取 Excel 文件。") from exc
    headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
    if headers != IMPORT_HEADERS:
        raise error(422, "IMPORT_TEMPLATE_INVALID", "模板列不匹配，请下载最新模板。")
    payload, lines = [], []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if row_number > settings.import_max_rows + 1:
            raise error(422, "IMPORT_TOO_MANY_ROWS", "导入行数超过限制。")
        values = dict(zip(IMPORT_HEADERS, row, strict=True))
        status, result, row_payload = import_row(values, db, user)
        payload.append(row_payload)
        lines.append((row_number, status, result))
    batch = StudentImportBatch(
        organization_id=user.organization_id,
        created_by=user.id,
        expires_at=utc_now() + timedelta(hours=2),
        preview_payload_ciphertext=encrypt(json.dumps(payload, ensure_ascii=False)),
        summary_json=json.dumps({"total": len(lines)}, ensure_ascii=False),
    )
    db.add(batch)
    db.flush()
    for row_number, status, result in lines:
        db.add(
            StudentImportLine(
                batch_id=batch.id,
                row_number=row_number,
                status=status,
                result_json=json.dumps(result, ensure_ascii=False),
            )
        )
    audit(
        db,
        user,
        request,
        "student.import.preview",
        "student_import_batch",
        batch.id,
        {"rows": len(lines)},
    )
    db.commit()
    return import_batch_view(db, batch)


def import_batch_view(db: Session, batch: StudentImportBatch) -> dict[str, Any]:
    lines = db.scalars(
        select(StudentImportLine)
        .where(StudentImportLine.batch_id == batch.id)
        .order_by(StudentImportLine.row_number)
    ).all()
    return {
        "id": batch.id,
        "status": batch.status,
        "expires_at": batch.expires_at,
        "items": [
            {
                "row_number": line.row_number,
                "status": line.status,
                "student_id": line.student_id,
                "result": json.loads(line.result_json),
            }
            for line in lines
        ],
    }


def import_batch_for_user(db: Session, user: User, batch_id: str) -> StudentImportBatch:
    batch = db.scalar(
        select(StudentImportBatch).where(
            StudentImportBatch.id == batch_id,
            StudentImportBatch.organization_id == user.organization_id,
            StudentImportBatch.created_by == user.id,
        )
    )
    if not batch:
        raise error(404, "IMPORT_BATCH_NOT_FOUND", "未找到导入预检批次。")
    return batch


@router.get("/imports/students/{batch_id}")
def get_import_batch(batch_id: str, db: Db, user: StudentImport) -> dict[str, Any]:
    return import_batch_view(db, import_batch_for_user(db, user, batch_id))


@router.post("/imports/students/{batch_id}/confirm")
def confirm_import(
    batch_id: str, data: ConfirmInput, request: Request, db: Db, user: StudentImport
) -> dict[str, Any]:
    batch = import_batch_for_user(db, user, batch_id)
    if batch.confirm_idempotency_key == data.idempotency_key:
        return import_batch_view(db, batch)
    if batch.status != "PREVIEWED":
        raise error(409, "IMPORT_BATCH_NOT_CONFIRMABLE", "预检批次不能再次确认。")
    if batch.expires_at < utc_now():
        raise error(409, "IMPORT_BATCH_EXPIRED", "预检批次已过期，请重新上传。")
    payloads = json.loads(decrypt(batch.preview_payload_ciphertext))
    lines = db.scalars(
        select(StudentImportLine)
        .where(StudentImportLine.batch_id == batch.id)
        .order_by(StudentImportLine.row_number)
    ).all()
    for index, line in enumerate(lines):
        item = payloads[index]
        if line.status == "DUPLICATE_CONFIRMED":
            continue
        if line.status == "POSSIBLE_DUPLICATE":
            action = data.possible_duplicate_actions.get(str(line.row_number), {})
            if action.get("action") == "skip":
                line.status = "SKIPPED"
                line.result_json = json.dumps(
                    {"message": "用户跳过疑似重复行。"}, ensure_ascii=False
                )
                continue
            if action.get("action") == "link":
                line.status = "MATCHED_EXISTING_STUDENT"
                line.student_id = action.get("student_id")
                continue
            if action.get("action") != "create" or not action.get("reason"):
                line.status = "INVALID"
                line.result_json = json.dumps(
                    {"message": "疑似重复行必须选择关联、强制新建并填写理由，或跳过。"},
                    ensure_ascii=False,
                )
                continue
            item["force_create_reason"] = action["reason"]
        if line.status not in {"READY_CREATE_STUDENT", "POSSIBLE_DUPLICATE"}:
            continue
        try:
            with db.begin_nested():
                student = create_student(db, user, request, StudentInput.model_validate(item))
                db.flush()
                line.status = "SUCCESS"
                line.student_id = student.id
                line.result_json = json.dumps(
                    {"student_no": student.student_no}, ensure_ascii=False
                )
        except HTTPException as exc:
            line.status = "FAILED"
            line.result_json = json.dumps(exc.detail, ensure_ascii=False)
    batch.status = "CONFIRMED"
    batch.confirm_idempotency_key = data.idempotency_key
    audit(db, user, request, "student.import.confirm", "student_import_batch", batch.id)
    db.commit()
    return import_batch_view(db, batch)


@router.get("/imports/students/{batch_id}/result")
def import_result(batch_id: str, db: Db, user: StudentImport) -> dict[str, Any]:
    return import_batch_view(db, import_batch_for_user(db, user, batch_id))
